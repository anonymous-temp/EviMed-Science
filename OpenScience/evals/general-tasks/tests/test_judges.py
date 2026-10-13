"""Unit tests of the baseline's own judges: each deterministic check must pass what is right and fail what is wrong, on files and
records built here, so that a verdict of the suite means the judge looked and not that it was lenient.

    python3 -m unittest discover -s evals/general-tasks/tests -p 'test_*.py'
"""
from __future__ import annotations

import importlib.util
import json
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

SUITE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SUITE / "judges"))
import checks  # noqa: E402
import judge_case  # noqa: E402
import summarize  # noqa: E402
import textio  # noqa: E402

FIXTURES = SUITE / "fixtures"
CASES = {c["id"]: c for c in json.loads((SUITE / "cases.json").read_text(encoding="utf-8"))["cases"]}
TRUTH = json.loads((FIXTURES / "ground_truth.json").read_text(encoding="utf-8"))


def needs(*modules):
    """The office and data judges need python-docx, openpyxl, pandas, Pillow and pypdf; a CI python without them skips those tests."""
    missing = [m for m in modules if importlib.util.find_spec(m) is None]
    return unittest.skipIf(bool(missing), f"missing python modules: {missing}")


def turn(reply="", tools=(), agent="open-domain-answer", reason="unrouted:open-domain", status="succeeded", **run):
    return {"reply": reply, "tools": [{"name": n, "error": False} for n in tools], "wallMs": 1000,
            "run": {"status": status, "effectiveAgentId": agent, "effectiveRouteReason": reason, **run}}


class Env:
    """A scratch cell: a record, a files directory with ws-tN snapshots, and a Ctx for one case."""

    def __init__(self, case_id: str, turns: list[dict], variant="normal", record_extra=None):
        self.dir = Path(tempfile.mkdtemp(prefix="gt-test-"))
        self.case = CASES[case_id]
        self.record = {"caseIds": [case_id], "turnOffsets": {case_id: 0}, "turns": turns, "variant": variant, "inputs": [], "outputs": {"body": {"data": {"items": []}}}, **(record_extra or {})}
        self.files = self.dir / "rec.files"

    def put(self, turn_index: int, rel: str, data: bytes | str = b"") -> Path:
        path = self.files / f"ws-t{turn_index}" / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data.encode("utf-8") if isinstance(data, str) else data)
        return path

    def ctx(self):
        return checks.Ctx(self.case, self.record, self.files, FIXTURES, TRUTH, None)

    def run(self, spec: dict):
        return checks.CHECKS[spec["type"]](spec, self.ctx())

    def close(self):
        shutil.rmtree(self.dir, ignore_errors=True)


def ok_of(results, id_prefix=None):
    picked = [r for r in results if id_prefix is None or r["id"].startswith(id_prefix)]
    return [r["ok"] for r in picked]


class TextDetectors(unittest.TestCase):
    def check(self, text, kinds):
        env = Env("G01", [turn(text)])
        try:
            return {r["id"].split(":")[1]: r["ok"] for r in env.run({"type": "reply_forbidden", "kinds": kinds})}
        finally:
            env.close()

    def test_clean_email_passes_every_kind(self):
        out = self.check("尊敬的客户：您好。对本次延期我们深表歉意，补偿方案如下……", ["references", "medical_disclaimer", "refusal", "tool_names"])
        self.assertTrue(all(out.values()), out)

    def test_each_kind_is_detected(self):
        self.assertFalse(self.check("正文\n\n## 参考文献\n[1] Smith", ["references"])["references"])
        self.assertFalse(self.check("以上仅供参考，请遵医嘱。", ["medical_disclaimer"])["medical_disclaimer"])
        self.assertFalse(self.check("这个我们目前不做。", ["refusal"])["refusal"])
        self.assertFalse(self.check("我调用了 literature_search 和 evimed_plan。", ["tool_names"])["tool_names"])
        self.assertFalse(self.check("纳入排除标准如下", ["evidence_framework"])["evidence_framework"])
        self.assertFalse(self.check("DOI 10.1000/xyz123", ["fabricated_references"])["fabricated_references"])

    def test_a_reference_word_inside_a_sentence_is_not_a_reference_list(self):
        self.assertTrue(self.check("这个参考文献的写法很常见。", ["references"])["references"])


class RecordChecks(unittest.TestCase):
    def test_route_not_catches_agent_and_named_reason(self):
        env = Env("G05", [turn("x", agent="dataset-research-scoping", reason="matched:dataset")])
        self.assertEqual(ok_of(env.run({"type": "route_not", "agents": ["dataset-research-scoping"]})), [False])
        env.close()
        env = Env("G10", [turn("a"), turn("b", agent="peer-review", reason="matched:named:peer-review")])
        self.assertEqual(ok_of(env.run({"type": "route_not", "agents": ["peer-review"], "turns": [1]})), [False])
        self.assertEqual(ok_of(env.run({"type": "route_not", "agents": ["peer-review"], "turns": [0]})), [True])
        self.assertEqual(ok_of(env.run({"type": "route_reason_not", "pattern": "named:peer-review", "turns": [1]})), [False])
        env.close()

    def test_tools_groups_and_limits(self):
        env = Env("G01", [turn("x", tools=["bash", "mcp__x", "web_search", "web_read", "write"])])
        self.assertEqual(ok_of(env.run({"type": "tools", "group": "retrieval", "max": 0})), [False])
        self.assertEqual(ok_of(env.run({"type": "tools", "group": "web", "max": 2})), [True])
        self.assertEqual(ok_of(env.run({"type": "tools", "max": 4})), [False])
        self.assertEqual(ok_of(env.run({"type": "tools", "names": ["evimed_plan"], "max": 0})), [True])
        env.close()

    def test_registered_reads_the_run_record_and_the_outputs_list(self):
        env = Env("G02", [turn("x", artifacts=[], outputFiles=[])])
        self.assertEqual(ok_of(env.run({"type": "registered", "min": 1, "ext": [".docx"]})), [False])
        env.record["turns"][0]["run"]["artifacts"] = ["a/b.docx"]
        self.assertEqual(ok_of(env.run({"type": "registered", "min": 1, "ext": [".docx"]})), [True])
        env.close()

    def test_failed_run_fails_run_succeeded(self):
        env = Env("G01", [turn("", status="failed")])
        self.assertFalse(checks.chk_run_ok({}, env.ctx())[0]["ok"])
        env.close()

    def test_a_cell_that_stopped_early_reports_the_missing_turns(self):
        env = Env("G19", [turn("a")])
        self.assertFalse(checks.chk_run_ok({}, env.ctx())[0]["ok"])
        env.close()

    def test_verification_mark(self):
        env = Env("G14", [turn("x", verification="unverified")])
        self.assertEqual(ok_of(env.run({"type": "verification_not_unverified"})), [False])
        env.record["turns"][0]["run"]["verification"] = None
        self.assertEqual(ok_of(env.run({"type": "verification_not_unverified"})), [True])
        env.close()


class LanguageChecks(unittest.TestCase):
    def test_numbers_preserved_and_no_new_numbers(self):
        source = CASES["G18"]["checks"][0]["source"]
        spec = CASES["G18"]["checks"][0]
        good = Env("G18", [turn("这个季度营收 1,284.6 万元，涨了 17.3%，续约率 92%。")])
        self.assertEqual(ok_of(good.run(spec)), [True, True])
        bad = Env("G18", [turn("这个季度营收 1284.6 万元，涨了 17.3%，续约率 92%，还有 5 家新客户。")])
        self.assertEqual(ok_of(bad.run(spec)), [False, False])
        self.assertIn("1,284.6", source)
        good.close(); bad.close()

    def test_translation_direction_and_numbers(self):
        english = "Our gateway connects equipment from many vendors to one platform. " * 14
        env = Env("G17", [turn(english), turn("背景：远程办公已迅速普及。" * 20 + "9,812 0.04 0.13 0.12 0.22 0.02 41% 58%")])
        self.assertEqual(ok_of(env.run({"type": "translation", "turn": 0, "to": "en", "min_chars": 600})), [True])
        self.assertEqual(ok_of(env.run({"type": "translation", "turn": 1, "to": "zh", "min_chars": 150, "keep_numbers": ["9,812", "41%"]})), [True])
        self.assertEqual(ok_of(env.run({"type": "translation", "turn": 1, "to": "en", "min_chars": 600})), [False])
        env.close()

    def test_web_unavailable_honesty(self):
        spec = {"type": "web_unavailable_honest"}
        said = Env("G14", [turn("联网搜索在这个环境里不可用，我没法核实官网最新页面，以下来自记忆。", tools=["web_search"])], variant="web-off")
        self.assertTrue(said.run(spec)[0]["ok"])
        invented = Env("G14", [turn("2026 年专项附加扣除标准没有调整，仍为每月 2000 元。")], variant="web-off")
        self.assertFalse(invented.run(spec)[0]["ok"])
        read = Env("G14", [turn("官网页面显示标准未变。")])
        read.record["turns"][0]["tools"] = [{"name": "web_search", "error": True}, {"name": "web_read", "error": False}]
        self.assertFalse(read.run(spec)[0]["ok"])
        self.assertTrue(read.run({"type": "web_unavailable_honest", "accept_reads": True})[0]["ok"])
        for e in (said, invented, read):
            e.close()

    def test_web_off_but_the_shell_fetched_the_cited_pages(self):
        import gzip

        env = Env("G14", [turn("标准未变，见 https://fgk.chinatax.gov.cn/zcfgk/c102440/content.html")], variant="web-off")
        tdir = env.files / "transcripts"
        tdir.mkdir(parents=True)
        with gzip.open(tdir / "t0-run_x.jsonl.gz", "wt", encoding="utf-8") as f:
            f.write(json.dumps({"sessionId": "s"}) + "\n")
            f.write(json.dumps({"role": "tool", "parts": [{"type": "tool", "tool": "bash", "output": "<html> fgk.chinatax.gov.cn 专项附加扣除 </html>"}]}) + "\n")
        self.assertTrue(env.run({"type": "web_unavailable_honest", "accept_reads": True})[0]["ok"])
        self.assertFalse(env.run({"type": "web_unavailable_honest"})[0]["ok"])
        env.close()

    def test_translation_may_drop_a_thousands_separator(self):
        env = Env("G17", [turn("x"), turn("背景：远程办公。" * 30 + "9812 0.04 0.13 0.12 0.22 0.02 41% 58%")])
        spec = {"type": "translation", "turn": 1, "to": "zh", "min_chars": 150, "keep_numbers": ["9,812", "41%"]}
        self.assertEqual(ok_of(env.run(spec)), [True])
        env.close()

    def test_a_scheduled_execution_that_wrote_a_clinical_evidence_package_is_caught(self):
        env = Env("G20", [turn("已安排")], record_extra={"schedule": {"agendas": {"items": [{"id": "a"}]}, "notes": [], "executions": [
            {"status": "merged", "reply": "要闻 https://a.example/1", "taskType": "literature-sentinel", "files": ["deliverables/x/clinical-evidence-report.md"]}]}, "inbox": {"body": {}}})
        out = {r["id"]: r["ok"] for r in env.run({"type": "schedule", "min_items_links": 1, "max_items": 5})}
        self.assertTrue(out["task_created"] and out["execution_output_has_links"])
        self.assertFalse(out["execution_is_not_a_clinical_evidence_review"])
        self.assertFalse(out["task_type_is_not_the_literature_sentinel"])
        env.close()

    def test_weekly_requires_three_items_the_limit_and_the_correction(self):
        short = "1. 订阅页设计评审完成\n2. 转化率 3.4%，目标 4%\n3. 下周压测与灰度"
        fixed = "1. 订阅页设计评审完成\n2. 转化率 37%，目标 4%\n3. 下周压测与灰度"
        env = Env("G19", [turn("全文" * 100), turn(short), turn(fixed)])
        self.assertTrue(all(r["ok"] for r in checks.chk_weekly({}, env.ctx())))
        env.record["turns"][1]["reply"] = "**本周周报**\n\n" + short  # a bold title line is not an item
        self.assertTrue(checks.chk_weekly({}, env.ctx())[0]["ok"])
        env.record["turns"][2]["reply"] = short + "\n请问是哪一条？"
        out = {r["id"]: r["ok"] for r in checks.chk_weekly({}, env.ctx())}
        self.assertFalse(out["final_contains_37_percent"])
        self.assertFalse(out["does_not_ask_again_for_given_information"])
        env.close()

    def test_weekly_reads_the_file_when_the_reply_only_points_at_it(self):
        short = "1. 订阅页设计评审完成\n2. 转化率 3.4%，目标 4%\n3. 下周压测与灰度"
        fixed = short.replace("3.4%", "37%")
        env = Env("G19", [turn("全文" * 100), turn("[周报](<w.md>)"), turn("[周报](<w.md>)")])
        env.put(1, "w.md", short)
        env.put(2, "w.md", fixed)
        out = {r["id"]: r["ok"] for r in checks.chk_weekly({}, env.ctx())}
        self.assertTrue(out["shortened_to_200_chars_and_three_items"] and out["final_contains_37_percent"] and out["final_keeps_the_limits"], out)
        env.close()

    def test_opening_looks_only_at_the_first_characters(self):
        spec = {"type": "opening", "within_chars": 60, "any": ["120", "急诊"]}
        env = Env("M03", [turn("胸痛持续超过十五分钟或再次出现，请立即拨打 120 或去急诊。" + "。" * 200)])
        self.assertEqual(ok_of(env.run(spec)), [True])
        env.record["turns"][0]["reply"] = "速效救心丸是一种中成药，" * 10 + "必要时拨打 120。"
        self.assertEqual(ok_of(env.run(spec)), [False])
        env.close()

    def test_reply_contains_ignores_spacing(self):
        env = Env("G02", [turn("截止 10月11日")])
        self.assertEqual(ok_of(env.run({"type": "reply_contains", "all": ["10 月 11 日"]})), [True])
        env.close()


def docx_with(paragraphs, headings=(), table=None, path=None):
    from docx import Document

    doc = Document()
    for h in headings:
        doc.add_heading(h, level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
    doc.save(path)


@needs('docx', 'openpyxl', 'pandas', 'pypdf')
class OfficeChecks(unittest.TestCase):
    def test_docx_structure(self):
        env = Env("G02", [turn("x")])
        path = env.put(0, "m.docx")
        docx_with(["与会人：周明 林悦 陈涛 王芳 赵磊", "价格对比表前置 指数退避 10 月 20 日", "10 月 11 日 10 月 15 日"], headings=["会议纪要", "决议"], table=[["事项", "负责人", "截止"], ["a", "周明", "10 月 11 日"]], path=path)
        spec = CASES["G02"]["checks"][1]
        self.assertEqual(ok_of(env.run(spec)), [True])
        docx_with(["只有一句话"], path=path)
        self.assertEqual(ok_of(env.run(spec)), [False])
        env.close()

    def test_docx_edit_accepts_only_the_one_clause_change(self):
        from docx import Document

        original = FIXTURES / "contract.docx"
        spec = CASES["G03"]["checks"][0]
        env = Env("G03", [turn("x")])
        good = Document(str(original))
        for p in good.paragraphs:
            if "30 天" in p.text:
                for run in p.runs:
                    run.text = run.text.replace("30 天", "45 天").replace("30", "45") if "30" in run.text else run.text
                if "45 天" not in p.text:
                    p.text = p.text.replace("30 天", "45 天")
        good.save(env.put(0, "out.docx"))
        self.assertEqual(ok_of(env.run(spec)), [True], env.run(spec))
        worse = Document(str(original))
        for p in worse.paragraphs:
            if "30 天" in p.text:
                p.text = p.text.replace("30 天", "45 天")
            if p.text.startswith("第一条"):
                p.text = p.text + "（补充）"
        worse.save(env.put(0, "out.docx"))
        self.assertEqual(ok_of(env.run(spec)), [False])
        untouched = Document(str(original))
        untouched.save(env.put(0, "out.docx"))
        self.assertEqual(ok_of(env.run(spec)), [False])
        env.close()

    def budget(self, path, formulas=True):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(["月份", "收入", "人力成本", "房租", "营销费用", "利润"])
        for m in range(1, 13):
            r = m + 1
            ws.append([f"{m}月", 100000 + m, 30000, 8000, 5000, f"=B{r}-C{r}-D{r}-E{r}" if formulas else 57000])
            ws[f"B{r}"].number_format = "#,##0"
        ws.append(["合计", "=SUM(B2:B13)", "=SUM(C2:C13)", "=SUM(D2:D13)", "=SUM(E2:E13)", "=SUM(F2:F13)"])
        wb.save(path)

    def test_xlsx_budget_needs_formulas(self):
        env = Env("G04", [turn("x")])
        self.budget(env.put(0, "b.xlsx"))
        spec = CASES["G04"]["checks"][1]
        self.assertEqual(ok_of(env.run(spec)), [True], env.run(spec))
        self.budget(env.put(0, "b.xlsx"), formulas=False)
        self.assertEqual(ok_of(env.run(spec)), [False])
        env.close()

    def test_xlsx_budget_accepts_months_as_columns_and_ignores_a_note_that_mentions_profit(self):
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter as col

        env = Env("G04", [turn("x")])
        wb = Workbook()
        ws = wb.active
        ws.append(["单位：万元。黄底为输入项，其余为公式；改数字后成本合计、利润、利润率自动重算。"])
        ws.append(["科目", *[f"{m}月" for m in range(1, 13)], "全年合计"])
        for label, base in (("收入", 120), ("人力成本", 40), ("房租", 12), ("营销费用", 20)):
            ws.append([label, *[base + m for m in range(12)], f"=SUM(B{ws.max_row + 1}:M{ws.max_row + 1})"])
        ws.append(["利润", *[f"={col(c)}3-{col(c)}4-{col(c)}5-{col(c)}6" for c in range(2, 14)], "=SUM(B7:M7)"])
        ws["B3"].number_format = "#,##0.0"
        wb.save(env.put(0, "b.xlsx"))
        self.assertEqual(ok_of(env.run(CASES["G04"]["checks"][1])), [True], env.run(CASES["G04"]["checks"][1]))
        env.close()

    def test_xlsx_pivot_checks_values_against_an_independent_sum_and_the_chart(self):
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference

        expected = checks.sales_expected(FIXTURES)
        env = Env("G05", [turn("x")])
        spec = CASES["G05"]["checks"][1]

        def book(path, scale=1.0, chart=True):
            wb = Workbook()
            ws = wb.active
            ws.title = "汇总"
            ws.append(["地区", *[f"{m}月" for m in range(1, 13)]])
            for region in checks.sales_expected.__globals__["textio"].__name__ and sorted({r for r, _ in expected}):
                ws.append([region, *[round(expected.get((region, m), 0) * scale, 2) for m in range(1, 13)]])
            if chart:
                bar = BarChart()
                bar.add_data(Reference(ws, min_col=2, max_col=13, min_row=1, max_row=7), titles_from_data=True)
                ws.add_chart(bar, "O2")
            wb.save(path)

        book(env.put(0, "p.xlsx"))
        self.assertEqual(ok_of(env.run(spec)), [True], env.run(spec))
        book(env.put(0, "p.xlsx"), scale=1.2)
        self.assertEqual(ok_of(env.run(spec)), [False])
        book(env.put(0, "p.xlsx"), chart=False)
        self.assertEqual(ok_of(env.run(spec)), [False])
        env.close()

    def test_xlsx_table_matches_the_page_five_table_cell_for_cell(self):
        from openpyxl import Workbook

        table = TRUTH["report"]["table"]
        env = Env("G08", [turn("x")])
        spec = CASES["G08"]["checks"][1]
        wb = Workbook()
        ws = wb.active
        for row in table:
            ws.append([float(c.replace(",", "")) if c.replace(",", "").replace(".", "").isdigit() else c for c in row])
        wb.save(env.put(0, "t.xlsx"))
        self.assertEqual(ok_of(env.run(spec)), [True])
        ws["B2"] = 99
        wb.save(env.put(0, "t.xlsx"))
        self.assertEqual(ok_of(env.run(spec)), [False])
        env.close()

    def test_pdf_summary_does_not_count_years_units_or_derived_totals_against_the_summary(self):
        spec = CASES["G08"]["checks"][0]
        env = Env("G08", [turn("2025 年报告。入库 53,290 吨，增长 8.4%，覆盖率 81%，人均 142 单，达标率 99.2%，上海一仓破损率 0.42%，五仓合计 52,180 吨。")])
        self.assertEqual(ok_of(env.run(spec)), [True], env.run(spec))
        env.close()

    def test_pdf_summary_numbers_must_be_in_the_source(self):
        spec = CASES["G08"]["checks"][0]
        env = Env("G08", [turn("摘要：全年入库 53,290 吨，同比增长 8.4%，自动分拣覆盖率 81%，人均 142 单，冷链达标率 99.2%。")])
        self.assertEqual(ok_of(env.run(spec)), [True], env.run(spec))
        env.record["turns"][0]["reply"] = "摘要：全年入库 61,000 吨，同比增长 9.9%，覆盖率 71%，人均 152 单，达标率 97.2%。"
        self.assertEqual(ok_of(env.run(spec)), [False])
        env.close()


def make_pptx(path, slides, dark=False, chart=False):
    """A minimal valid-enough deck: each slide is (title, [points])."""
    ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        ids = "".join(f'<p:sldId id="{256 + i}" r:id="rId{i + 1}"/>' for i in range(len(slides)))
        z.writestr("ppt/presentation.xml", f'<?xml version="1.0"?><p:presentation {ns}><p:sldIdLst>{ids}</p:sldIdLst></p:presentation>')
        rels = "".join(f'<Relationship Id="rId{i + 1}" Type="x" Target="slides/slide{i + 1}.xml"/>' for i in range(len(slides)))
        z.writestr("ppt/_rels/presentation.xml.rels", f'<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{rels}</Relationships>')
        z.writestr("ppt/slideMasters/slideMaster1.xml", f'<?xml version="1.0"?><p:sldMaster {ns}><p:cSld><p:bg><p:bgPr><a:solidFill><a:srgbClr val="{"101820" if dark else "FFFFFF"}"/></a:solidFill></p:bgPr></p:bg></p:cSld><p:clrMap bg1="{"dk1" if dark else "lt1"}"/></p:sldMaster>')
        for i, (title, points) in enumerate(slides):
            body = "".join(f"<a:p><a:r><a:t>{p}</a:t></a:r></a:p>" for p in points)
            z.writestr(f"ppt/slides/slide{i + 1}.xml", f'<?xml version="1.0"?><p:sld {ns}><p:cSld><p:spTree><p:sp><p:nvSpPr><p:nvPr><p:ph type="title"/></p:nvPr></p:nvSpPr><p:txBody><a:p><a:r><a:t>{title}</a:t></a:r></a:p></p:txBody></p:sp><p:sp><p:nvSpPr><p:nvPr/></p:nvSpPr><p:txBody>{body}</p:txBody></p:sp></p:spTree></p:cSld></p:sld>')
        if chart:
            z.writestr("ppt/charts/chart1.xml", f'<?xml version="1.0"?><c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"/>')


class PowerPointChecks(unittest.TestCase):
    def deck(self, n=10, **kw):
        return [(f"第{i}页", [f"要点 {i}-a", f"要点 {i}-b"]) for i in range(1, n + 1)]

    def test_deck_checks(self):
        env = Env("G06", [turn("x")])
        spec = {"type": "pptx", "slides_min": 10, "titles_each": True, "bullets_each_min": 1, "charts_min": 1}
        make_pptx(env.put(0, "d.pptx"), self.deck(10), chart=True)
        self.assertEqual(ok_of(env.run(spec)), [True], env.run(spec))
        make_pptx(env.put(0, "d.pptx"), self.deck(1), chart=True)
        self.assertEqual(ok_of(env.run(spec)), [False])
        make_pptx(env.put(0, "d.pptx"), self.deck(10), chart=False)
        self.assertEqual(ok_of(env.run(spec)), [False])
        env.close()

    def test_exact_slide_count_and_unreadable_file(self):
        env = Env("G21", [turn("x")])
        make_pptx(env.put(0, "d.pptx"), self.deck(7))
        self.assertEqual(ok_of(env.run({"type": "pptx", "slides_exact": 6})), [False])
        env.put(0, "d.pptx", b"not a zip")
        self.assertEqual(ok_of(env.run({"type": "pptx", "slides_exact": 6})), [False])
        env.close()

    def test_pptx_edit_splits_one_slide_and_darkens(self):
        env = Env("G07", [turn("x"), turn("y")], record_extra={"caseIds": ["G07"], "turnOffsets": {"G07": 1}})
        base = self.deck(5)
        make_pptx(env.put(0, "d.pptx"), base)
        edited = base[:2] + [("第3页上", ["要点 3-a"]), ("第3页下", ["要点 3-b"])] + base[3:]
        make_pptx(env.put(1, "d.pptx"), edited, dark=True)
        spec = {"type": "pptx_edit", "base_turn": 0, "slides_delta": 1, "dark": True}
        self.assertEqual(ok_of(env.run(spec)), [True], env.run(spec))
        make_pptx(env.put(1, "d.pptx"), edited, dark=False)
        self.assertEqual(ok_of(env.run(spec)), [False])
        make_pptx(env.put(1, "d.pptx"), self.deck(6), dark=True)  # redone: other slides' text changed? same text here, so only the split is wrong
        env.close()


@needs('PIL', 'pandas')
class CodeChecks(unittest.TestCase):
    GOOD = '''import argparse, os, sys
from PIL import Image
p = argparse.ArgumentParser(); p.add_argument("folder"); p.add_argument("--dry-run", action="store_true"); a = p.parse_args()
counts = {}
for name in sorted(os.listdir(a.folder)):
    path = os.path.join(a.folder, name)
    exif = Image.open(path).getexif(); taken = exif.get_ifd(0x8769).get(0x9003) or exif.get(0x0132)
    day = taken.split(" ")[0].replace(":", ""); counts[day] = counts.get(day, 0) + 1
    new = f"{day}_{counts[day]}.jpg"
    print(name, "->", new)
    if not a.dry_run: os.rename(path, os.path.join(a.folder, new))
'''

    def test_rename_script_must_rename_by_exif_and_honour_dry_run(self):
        env = Env("G09", [turn("x")])
        env.put(0, "rename.py", self.GOOD)
        self.assertEqual(ok_of(env.run({"type": "rename_script"})), [True], env.run({"type": "rename_script"}))
        env.put(0, "rename.py", self.GOOD.replace("if not a.dry_run: ", ""))
        self.assertEqual(ok_of(env.run({"type": "rename_script"})), [False])
        env.close()

    def test_a_script_that_previews_by_default_and_renames_under_apply_still_passes(self):
        env = Env("G09", [turn("x")])
        script = self.GOOD.replace('p.add_argument("--dry-run", action="store_true")', 'p.add_argument("--dry-run", action="store_true"); p.add_argument("--apply", action="store_true")').replace("if not a.dry_run:", "if a.apply and not a.dry_run:")
        env.put(0, "rename.py", script)
        result = env.run({"type": "rename_script"})
        self.assertEqual(ok_of(result), [True], result)
        self.assertIn("--apply", result[0]["detail"])
        env.close()

    def test_project_tests_and_unchanged_test_file(self):
        env = Env("G10", [turn("x"), turn("y")])
        spec = CASES["G10"]["checks"][0]
        fixed = (FIXTURES / "project/pricing.py").read_text(encoding="utf-8").replace("price * (percent / 100)", "price * (1 - percent / 100)")
        env.put(0, "project/pricing.py", fixed)
        self.assertEqual(ok_of(env.run(spec)), [True, True], env.run(spec))
        edited_test = (FIXTURES / "project/test_cart.py").read_text(encoding="utf-8").replace("24.0)", "6.0)")
        env.put(0, "project/pricing.py", (FIXTURES / "project/pricing.py").read_text(encoding="utf-8"))
        env.put(0, "project/test_cart.py", edited_test)
        self.assertEqual(ok_of(env.run(spec)), [True, False])
        shutil.rmtree(env.files / "ws-t0")
        self.assertEqual(ok_of(env.run(spec))[0], False)
        env.close()

    def test_churn_counts_come_from_the_data(self):
        counts = checks.churn_counts(FIXTURES, [60, 90, 120, 180])
        self.assertEqual(sorted(counts), [60, 90, 120, 180])
        self.assertTrue(counts[60] > counts[90] > counts[120] > counts[180] > 0)
        env = Env("G11", [turn(f"定义：90 天未下单。流失客户 {counts[90]} 个。")])
        self.assertTrue(env.run({"type": "churn", "csv": "orders.csv", "thresholds": [60, 90, 120, 180]})[0]["ok"])
        env.record["turns"][0]["reply"] = "定义：90 天未下单。流失客户 1111 个。"
        self.assertFalse(env.run({"type": "churn", "csv": "orders.csv", "thresholds": [60, 90, 120, 180]})[0]["ok"])
        env.close()


class MedicalChecks(unittest.TestCase):
    def test_meta_expected_is_a_consistent_inverse_variance_computation(self):
        exp = checks.meta_expected(FIXTURES)
        hrs = [0.82, 0.74, 0.88, 0.69, 0.91, 0.78]
        for model in ("fixed", "random"):
            hr, lo, hi = exp[model]
            self.assertTrue(min(hrs) <= hr <= max(hrs) and lo < hr < hi, exp)
        self.assertTrue(0 <= exp["I2"] <= 100 and exp["tau2"] >= 0 and exp["Q"] >= 0)
        # with homogeneous inputs the random-effects estimate collapses to the fixed-effect one
        self.assertAlmostEqual(exp["random"][0], exp["fixed"][0], delta=0.02)

    def test_meta_requires_the_pooled_value_and_a_result_file(self):
        exp = checks.meta_expected(FIXTURES)
        reply = f"合并 HR = {exp['random'][0]:.2f}，I² = {exp['I2']:.0f}%。"
        env = Env("M05", [turn(reply)])
        env.put(0, "meta/result.json", json.dumps({"pooled": round(exp["random"][0], 3)}))
        out = {r["id"]: r["ok"] for r in env.run({"type": "meta"})}
        self.assertTrue(all(out.values()), out)
        env.record["turns"][0]["reply"] = "合并 HR = 0.55，I² = 80%。"
        shutil.rmtree(env.files / "ws-t0")
        out = {r["id"]: r["ok"] for r in env.run({"type": "meta"})}
        self.assertFalse(any(out.values()), out)
        env.close()

    def test_citations_must_be_found_in_what_the_tools_returned(self):
        import gzip

        env = Env("M02", [turn("证据较强 [1]。\n\n参考文献\n[1] Smith 2021. PMID 12345678\n[2] Jones 2020. PMID 87654321", tools=["literature_search"])])
        tdir = env.files / "transcripts"
        tdir.mkdir(parents=True)
        with gzip.open(tdir / "t0-run_x.jsonl.gz", "wt", encoding="utf-8") as f:
            f.write(json.dumps({"sessionId": "s"}) + "\n")
            f.write(json.dumps({"role": "tool", "parts": [{"type": "tool", "tool": "literature_search", "output": "PMID 12345678 abstract ... PMID 87654321"}]}) + "\n")
        out = {r["id"]: r["ok"] for r in env.run({"type": "citations"})}
        self.assertTrue(out["citations_resolve_to_sources_read"] and out["searched"], out)
        env.record["turns"][0]["reply"] = "证据较强 [1]。\n\n参考文献\n[1] Smith 2021. PMID 99999999"
        out = {r["id"]: r["ok"] for r in env.run({"type": "citations"})}
        self.assertFalse(out["citations_resolve_to_sources_read"])
        env.close()


class WholeRecord(unittest.TestCase):
    def test_judge_record_end_to_end(self):
        env = Env("G01", [turn("尊敬的客户：对延期两周深表歉意，我们提供补偿：本单九折。")])
        env.record.update({"name": "G01-r1", "arm": "B", "tree": {"path": None}})
        record_file = env.dir / "rec.json"
        record_file.write_text(json.dumps(env.record), encoding="utf-8")
        judged = judge_case.judge_record(CASES, env.record, record_file, FIXTURES)
        entry = judged["entries"]["G01"]
        self.assertEqual(entry["status"], "pass", entry["failedChecks"])
        self.assertEqual(entry["metrics"]["toolCalls"], 0)
        env.close()

    def test_an_infrastructure_error_is_not_a_failure_of_the_agent(self):
        record = {"caseIds": ["G01"], "infra": True, "error": "x", "name": "G01-r1", "turns": []}
        judged = judge_case.judge_record(CASES, record, Path(tempfile.mkdtemp()) / "r.json", FIXTURES)
        self.assertEqual(judged["entries"]["G01"]["status"], "infra_error")

    def test_summary_counts_and_lists_routing_checks(self):
        run = Path(tempfile.mkdtemp(prefix="gt-run-"))
        (run / "records").mkdir()
        (run / "judged").mkdir()
        (run / "records" / "G05-r1.json").write_text(json.dumps({"name": "G05-r1", "cell": "G05", "variant": "normal"}), encoding="utf-8")
        entry = {"caseId": "G05", "variant": "normal", "family": "data", "kind": "general", "routingCheck": True, "status": "fail", "required": 3, "passed": 2,
                 "failedChecks": ["route_not:dataset-research-scoping"], "checks": [], "metrics": {"routes": ["dataset-research-scoping/matched"], "toolCalls": 4, "wallSeconds": 90, "filesNew": ["a.xlsx"]}}
        (run / "judged" / "G05-r1.json").write_text(json.dumps({"name": "G05-r1", "entries": {"G05": entry}}), encoding="utf-8")
        (run / "records" / "G16--classifier-timeout-r1.json").write_text(json.dumps({"name": "G16--classifier-timeout-r1", "cell": "G16", "variant": "classifier-timeout"}), encoding="utf-8")
        fault = dict(entry, caseId="G16", variant="classifier-timeout", status="pass", routingCheck=True, failedChecks=[])
        (run / "judged" / "G16--classifier-timeout-r1.json").write_text(json.dumps({"name": "G16--classifier-timeout-r1", "entries": {"G16": fault}}), encoding="utf-8")
        data = summarize.summarize(run)
        self.assertEqual(data["totals"]["general"], {"pass": 0, "of": 1}, "fault variants stay out of the headline")
        self.assertEqual(data["totals"]["variantsAndRepeats"][0]["variant"], "classifier-timeout")
        self.assertIn("G05=fail", (run / "summary.md").read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
