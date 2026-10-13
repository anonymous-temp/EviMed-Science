"""Readers for the files the judges look at: Word, PowerPoint, Excel, PDF and plain text.

Standard library plus python-docx, openpyxl and pypdf (all already used by this repository's other harnesses). PowerPoint
is read straight from its zip parts so that no python-pptx is needed.
"""
from __future__ import annotations

import re
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def read_text_file(path: Path) -> str:
    return Path(path).read_text(encoding="utf-8", errors="replace")


def strip_html(text: str) -> str:
    text = re.sub(r"(?is)<(script|style).*?</\1>", " ", text)
    return re.sub(r"\s+", " ", re.sub(r"(?s)<[^>]+>", " ", text))


def docx_open(path: Path):
    from docx import Document

    return Document(str(path))


def docx_text(path: Path) -> str:
    doc = docx_open(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    for section in doc.sections:
        parts += [p.text for p in section.header.paragraphs] + [p.text for p in section.footer.paragraphs]
    return "\n".join(parts)


def docx_tables(path: Path) -> list[list[list[str]]]:
    doc = docx_open(path)
    return [[[cell.text.strip() for cell in row.cells] for row in table.rows] for table in doc.tables]


def pdf_pages(path: Path) -> list[str]:
    from pypdf import PdfReader

    return [(page.extract_text() or "") for page in PdfReader(str(path)).pages]


def pdf_text(path: Path) -> str:
    return "\n".join(pdf_pages(path))


def _slide_key(name: str) -> int:
    match = re.search(r"slide(\d+)\.xml$", name)
    return int(match.group(1)) if match else 0


def pptx_info(path: Path) -> dict:
    """Slides in presentation order with their title, texts, charts and pictures; raises if the zip or an XML part is broken."""
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        for name in names:
            if name.endswith(".xml") or name.endswith(".rels"):
                ET.fromstring(archive.read(name))  # every XML part must parse
        order = _slide_order(archive)
        slides = []
        for name in order:
            root = ET.fromstring(archive.read(name))
            rels_name = name.replace("slides/", "slides/_rels/") + ".rels"
            rels = archive.read(rels_name).decode("utf-8", "replace") if rels_name in names else ""
            title = ""
            texts = []
            for shape in root.iter("{%s}sp" % NS["p"]):
                placeholder = shape.find(".//p:nvPr/p:ph", NS)
                paragraphs = ["".join(t.text or "" for t in para.iter("{%s}t" % NS["a"])) for para in shape.iter("{%s}p" % NS["a"])]
                paragraphs = [p for p in paragraphs if p.strip()]
                if placeholder is not None and placeholder.get("type") in ("title", "ctrTitle"):
                    title = " ".join(paragraphs)
                else:
                    texts += paragraphs
            tables = len(list(root.iter("{%s}tbl" % NS["a"])))
            background = _background_rgb(root)
            slides.append({
                "part": name, "title": title, "texts": texts, "charts": rels.count("/chart"), "pictures": len(list(root.iter("{%s}pic" % NS["p"]))),
                "tables": tables, "background": background,
                "all_text": " ".join(t.text or "" for t in root.iter("{%s}t" % NS["a"])),
            })
        master_bg = _master_background(archive)
        theme_dark = _theme_is_dark(archive)
        chart_parts = [n for n in names if re.match(r"ppt/charts/chart\d+\.xml$", n)]
        media = [n for n in names if n.startswith("ppt/media/")]
    return {"slides": slides, "chart_parts": len(chart_parts), "media": len(media), "master_background": master_bg, "theme_dark": theme_dark}


def _slide_order(archive: zipfile.ZipFile) -> list[str]:
    names = archive.namelist()
    try:
        presentation = ET.fromstring(archive.read("ppt/presentation.xml"))
        rels = ET.fromstring(archive.read("ppt/_rels/presentation.xml.rels"))
        targets = {rel.get("Id"): rel.get("Target") for rel in rels}
        ordered = []
        for slide_id in presentation.iter("{%s}sldId" % NS["p"]):
            target = targets.get(slide_id.get("{%s}id" % NS["r"]))
            if target:
                ordered.append("ppt/" + target.lstrip("/").replace("ppt/", "", 1) if not target.startswith("/") else target.lstrip("/"))
        ordered = [n for n in ordered if n in names]
        if ordered:
            return ordered
    except (KeyError, ET.ParseError):
        pass
    return sorted((n for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n)), key=_slide_key)


def _rgb_of(element) -> str | None:
    if element is None:
        return None
    srgb = element.find(".//a:srgbClr", NS)
    if srgb is not None and srgb.get("val"):
        return srgb.get("val")
    return None


def _background_rgb(root) -> str | None:
    bg = root.find(".//p:bg", NS)
    return _rgb_of(bg)


def _master_background(archive: zipfile.ZipFile) -> str | None:
    for name in archive.namelist():
        if re.match(r"ppt/slideMasters/slideMaster\d+\.xml$", name):
            return _background_rgb(ET.fromstring(archive.read(name)))
    return None


def _theme_is_dark(archive: zipfile.ZipFile) -> bool | None:
    """Whether the master maps background 1 to a dark theme colour (the scheme's dk1)."""
    for name in archive.namelist():
        if re.match(r"ppt/slideMasters/slideMaster\d+\.xml$", name):
            root = ET.fromstring(archive.read(name))
            mapping = root.find(".//p:clrMap", NS)
            if mapping is not None:
                return mapping.get("bg1") == "dk1"
    return None


def luminance(rgb: str) -> float:
    r, g, b = (int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def pptx_text(path: Path) -> str:
    info = pptx_info(path)
    return "\n".join(slide["all_text"] for slide in info["slides"])


def xlsx_load(path: Path, data_only: bool = False):
    from openpyxl import load_workbook

    return load_workbook(str(path), data_only=data_only)


def xlsx_text(path: Path) -> str:
    wb = xlsx_load(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            out.append(" ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


class NoSpreadsheetEngine(RuntimeError):
    """LibreOffice Calc is not installed (or cannot load the file), so formula cells cannot be given values."""


def recalculated(path: Path) -> Path:
    """A copy of a workbook with every formula's cached value filled in by LibreOffice Calc (openpyxl writes formulas without values)."""
    out_dir = Path(tempfile.mkdtemp(prefix="gt-recalc-"))
    profile = Path(tempfile.mkdtemp(prefix="gt-lo-"))
    try:
        subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out_dir), str(path)],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)
    except (OSError, subprocess.SubprocessError) as error:
        raise NoSpreadsheetEngine(f"soffice failed: {type(error).__name__}") from error
    made = list(out_dir.glob("*.xlsx"))
    if not made:
        raise NoSpreadsheetEngine("LibreOffice Calc produced no workbook (is libreoffice-calc installed?)")
    return made[0]


def to_pdf(path: Path) -> Path:
    """Convert an Office file to PDF with LibreOffice (the stand-in for 'it opens'); raises if LibreOffice cannot."""
    out_dir = Path(tempfile.mkdtemp(prefix="gt-pdf-"))
    profile = Path(tempfile.mkdtemp(prefix="gt-lo-"))
    subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(path)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    return next(out_dir.glob("*.pdf"))


def file_text(path: Path) -> str:
    suffix = Path(path).suffix.lower()
    try:
        if suffix == ".docx":
            return docx_text(path)
        if suffix == ".pptx":
            return pptx_text(path)
        if suffix in (".xlsx", ".xlsm"):
            return xlsx_text(path)
        if suffix == ".pdf":
            return pdf_text(path)
        if suffix in (".html", ".htm"):
            return strip_html(read_text_file(path))
        if suffix in (".md", ".txt", ".csv", ".json", ".py", ".ts", ".js", ".tsv"):
            return read_text_file(path)
    except Exception:  # an unreadable file contributes no text; the opens-check reports it separately
        return ""
    return ""
