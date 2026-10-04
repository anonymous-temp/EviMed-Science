#!/usr/bin/env python3
"""The raw measurements behind a source's structured materials, taken inside a
disposable container: a PDF's text, page by page, and a spreadsheet's cells.

The in-house parser sends no pages, no regions and no cell spans (see
``packages/domain/src/sourceMaterials.mjs``). The platform holds the original
bytes, so what the parser cannot say is measured here from them: the control
plane stages ONE copy of the source file in a scratch directory it shares with
the runtime controller and asks the controller to run this script in a bounded,
network-less container (``apps/server/src/vcrIntakeController.mjs``, operation
``materials``, the same mechanism as the figure digitizer). The container sees
exactly two paths -- ``/input`` (the request and the one file, read-only) and
``/output`` (empty) -- and nothing else: no model, no workspace, no credential.

This script measures; the control plane decides. It reports each page's text and
whether the page has a text layer, and each sheet's cells with their addresses,
and never says which table sits on which page or what a number means: that rule
lives in one place (``sourceMaterialsLocate.mjs``) and is tested without Python.

- PDF: ``pypdf``'s text layer, page by page, into ``pages.json`` (a list of page
  texts, page 1 first) with ``result.json`` carrying the counts. A page with no
  text layer is a page with no text -- which is how a scan reads -- and stays in
  the count.
- Spreadsheet (``.xlsx`` / ``.xlsm``): ``openpyxl``'s view of every sheet, into
  ``result.json``: the cell's address, its type (number, string, date, boolean,
  error), its value, its formula (when it has one) and the cached value of that
  formula only when the file holds one -- a formula nobody computed has no value
  here, and is reported with none. Merged ranges and the sheet state (visible or
  hidden) are reported. At most ``maxCellsPerSheet`` cells of a sheet are listed;
  ``totalCells`` says how many it has.

Hidden knowledge: a ``.xlsx`` is a zip and its XML is untrusted, so before
openpyxl sees it the member count and the declared unpacked size are bounded,
an encrypted member is refused, and no XML part may carry a DOCTYPE or an ENTITY
(Excel never writes one). None of this protects the host -- the container does --
it makes the refusal a reason a reader can act on instead of a killed process.
The class of an exception, never its message, reaches the control plane: the
message may quote the document.
"""

from __future__ import annotations

import argparse
import datetime
import io
import json
import re
import signal
import sys
import zipfile
from pathlib import Path

import vcr_record_extract as record

NAME = "evimed-source-material-extract"
VERSION = "1.0.0"
PROTOCOL = 1

# A page with fewer significant characters than this has no text layer: a page
# number and a running head do not make a scanned page a text page.
DEFAULT_MIN_PAGE_CHARS = 20
MAX_ZIP_MEMBERS = 4000
MAX_ZIP_UNPACKED_BYTES = 256 * 1024 * 1024
MAX_XML_SCAN_BYTES = 24 * 1024 * 1024
# A sheet whose bounding box holds more cells than this is listed by its size and not read.
MAX_SHEET_CELLS = 2_000_000
MAX_CELL_TEXT = 1000


class Refusal(Exception):
    """A document this script will not read, with the reason the control plane maps."""

    def __init__(self, reason: str, detail: str = ""):
        super().__init__(reason)
        self.reason = reason
        self.detail = detail


def _libraries() -> dict:
    versions: dict[str, str] = {}
    for module in ("pypdf", "openpyxl"):
        try:
            imported = __import__(module)
            versions[module] = str(getattr(imported, "__version__", "unknown"))
        except ImportError:
            versions[module] = "missing"
    versions["python"] = "%d.%d.%d" % sys.version_info[:3]
    return versions


def _extractor() -> dict:
    return {"name": NAME, "version": VERSION, "libraries": _libraries()}


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def read_pdf_pages(data: bytes, max_pages: int, max_chars: int, min_page_chars: int) -> dict:
    """The text of every page, and whether each has a text layer."""
    try:
        from pypdf import PdfReader  # imported late: the module is also imported where pypdf is absent
    except ImportError as error:  # pragma: no cover - the image ships pypdf
        raise Refusal("converter_missing", "pypdf is not installed") from error
    if b"%PDF-" not in data[:1024]:
        raise Refusal("not_pdf")
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted and not reader.decrypt(""):
            raise Refusal("encrypted")
        pages = len(reader.pages)
    except Refusal:
        raise
    except Exception as error:  # noqa: BLE001 - pypdf raises a wide family; the reason is what matters
        name = type(error).__name__
        if "Dependency" in name or "encrypt" in str(error).lower() or "password" in str(error).lower():
            raise Refusal("encrypted", name) from error
        raise Refusal("corrupt", name) from error
    if pages < 1:
        raise Refusal("corrupt", "no pages")
    if pages > max_pages:
        raise Refusal("too_many_pages", str(pages))
    texts: list[str] = []
    layers: list[bool] = []
    chars: list[int] = []
    page_errors = 0
    total = 0
    truncated = False
    for index in range(pages):
        try:
            extracted = reader.pages[index].extract_text() or ""
        except Exception:  # noqa: BLE001 - one unreadable page is a page without text, not a failed document
            extracted = ""
            page_errors += 1
        cleaned = record.clean_text(extracted)
        total += len(cleaned)
        if total > max_chars:
            truncated = True
            break
        significant = record.significant_chars(cleaned)
        texts.append(cleaned)
        chars.append(significant)
        layers.append(significant >= min_page_chars)
    return {
        "pages": pages, "pagesRead": len(texts), "pageChars": chars, "textLayer": layers,
        "texts": texts, "truncated": truncated, "pageErrors": page_errors,
    }


# ---------------------------------------------------------------------------
# Spreadsheets
# ---------------------------------------------------------------------------


def _check_package(data: bytes) -> None:
    """Refuse what no spreadsheet is: not a zip, encrypted, absurdly large or carrying a DTD."""
    if data[:4] != b"PK\x03\x04":
        # An OLE compound file is an old .xls, or an encrypted package: neither is read here.
        raise Refusal("not_xlsx")
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as error:
        raise Refusal("corrupt", "zip") from error
    with archive:
        members = archive.infolist()
        if len(members) > MAX_ZIP_MEMBERS:
            raise Refusal("corrupt", "members")
        if sum(member.file_size for member in members) > MAX_ZIP_UNPACKED_BYTES:
            raise Refusal("corrupt", "unpacked size")
        if any(member.flag_bits & 0x1 for member in members):
            raise Refusal("encrypted")
        names = {member.filename for member in members}
        if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
            raise Refusal("not_xlsx", "no workbook part")
        for member in members:
            if not member.filename.endswith((".xml", ".rels")):
                continue
            if member.file_size > MAX_XML_SCAN_BYTES:
                raise Refusal("too_large", "xml part")
            with archive.open(member) as stream:
                head = stream.read(MAX_XML_SCAN_BYTES + 1)
            if len(head) > MAX_XML_SCAN_BYTES:
                raise Refusal("too_large", "xml part")
            if re.search(rb"<!(?:doctype|entity)", head, re.IGNORECASE):
                raise Refusal("corrupt", "dtd")


def _plain(value: object):
    """A cell value as JSON: numbers, strings and booleans as they are, dates as ISO text."""
    if isinstance(value, bool):
        return "b", value
    if isinstance(value, (int, float)):
        return "n", value
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return "d", value.isoformat()
    text = getattr(value, "text", None)  # an array formula object carries its formula text
    if isinstance(text, str):
        return "s", text
    return "s", str(value)


def read_xlsx_sheets(data: bytes, max_sheets: int, max_cells_per_sheet: int) -> dict:
    """Every sheet's cells with their addresses."""
    _check_package(data)
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - the image ships openpyxl
        raise Refusal("converter_missing", "openpyxl is not installed") from error
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    except Exception as error:  # noqa: BLE001 - the class is the reason; the message may quote the file
        raise Refusal("corrupt", type(error).__name__) from error
    cached = None
    sheets: list[dict] = []
    for position, sheet in enumerate(workbook.worksheets):
        if position >= max_sheets:
            break
        rows = int(sheet.max_row or 0)
        columns = int(sheet.max_column or 0)
        entry: dict = {
            "name": str(sheet.title)[:120], "index": position + 1, "state": str(sheet.sheet_state),
            "dimensions": {"rows": rows, "cols": columns},
            "merges": [str(item) for item in list(sheet.merged_cells.ranges)[:200]],
            "cells": [], "totalCells": 0, "truncated": False, "beyond": 0,
        }
        if rows * columns > MAX_SHEET_CELLS:
            entry["skipped"] = "too_large"
            sheets.append(entry)
            continue
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                entry["totalCells"] += 1
                if len(entry["cells"]) >= max_cells_per_sheet:
                    entry["truncated"] = True
                    # Cells that start like a number and are not listed: counted, so the control plane can say how many values it left unread.
                    if isinstance(value, (int, float)) and not isinstance(value, bool) or (isinstance(value, str) and re.match(r"^[(\[<>+\-~]?\s*\.?\d", value)):
                        entry["beyond"] += 1
                    continue
                formula = None
                if cell.data_type == "f":
                    formula = str(value if isinstance(value, str) else getattr(value, "text", value))[:500]
                    if cached is None:
                        try:
                            cached = load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
                        except Exception:  # noqa: BLE001 - no cached values is a fact about the file
                            cached = False
                    value = cached[sheet.title][cell.coordinate].value if cached else None
                if value is None:
                    kind, plain = "n", None  # a formula with no cached value: no value is reported
                else:
                    kind, plain = _plain(value)
                    if kind == "s":
                        plain = plain[:MAX_CELL_TEXT]
                        if cell.data_type == "e":
                            kind = "e"
                item = {"a": cell.coordinate, "k": kind, "v": plain}
                if formula is not None:
                    item["f"] = formula
                if cell.number_format and cell.number_format != "General":
                    item["nf"] = str(cell.number_format)[:40]
                entry["cells"].append(item)
        sheets.append(entry)
    return {"sheets": sheets, "sheetsTotal": len(workbook.worksheets), "truncated": len(workbook.worksheets) > max_sheets}


# ---------------------------------------------------------------------------
# The container's entry
# ---------------------------------------------------------------------------


def run(request: dict, input_dir: Path, output_dir: Path) -> dict:
    """One measurement: the request names the staged file, its digest and size, the format and the limits."""
    spec = request.get("file") or {}
    name = str(spec.get("name") or "")
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,63}", name):
        raise Refusal("request_invalid", "name")
    try:
        data = record.read_verified(input_dir / name, {"sha256": spec.get("sha256"), "bytes": spec.get("bytes")})
    except record.Refusal as refusal:
        raise Refusal(refusal.reason, refusal.detail) from refusal
    form = str(request.get("format") or "")
    limits = request.get("limits") or {}
    if form == "pdf":
        got = read_pdf_pages(
            data, int(limits.get("maxPages", 1000)), int(limits.get("maxChars", 16 * 1024 * 1024)), int(limits.get("minPageChars", DEFAULT_MIN_PAGE_CHARS)),
        )
        texts = got.pop("texts")
        (output_dir / "pages.json").write_text(json.dumps(texts, ensure_ascii=False), encoding="utf-8")
        return {"protocol": PROTOCOL, "outcome": "pages", "extractor": _extractor(), "format": "pdf", **got}
    if form in ("xlsx", "xlsm"):
        got = read_xlsx_sheets(data, int(limits.get("maxSheets", 40)), int(limits.get("maxCellsPerSheet", 1000)))
        return {"protocol": PROTOCOL, "outcome": "cells", "extractor": _extractor(), "format": form, **got}
    raise Refusal("request_invalid", "format")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--request", required=True)
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--deadline", type=int, default=0)
    args = parser.parse_args(argv)
    output_dir = Path(args.output_dir)
    base = {"protocol": PROTOCOL, "extractor": _extractor()}

    def on_alarm(_signal: int, _frame: object) -> None:
        raise Refusal("deadline")

    if args.deadline > 0 and hasattr(signal, "SIGALRM"):
        signal.signal(signal.SIGALRM, on_alarm)
        signal.alarm(args.deadline)
    try:
        request = json.loads(Path(args.request).read_text(encoding="utf-8"))
        result = run(request, Path(args.input_dir), output_dir)
    except Refusal as refusal:
        result = {**base, "outcome": "refused", "reason": refusal.reason}
    except MemoryError:
        result = {**base, "outcome": "refused", "reason": "memory"}
    except Exception as error:  # noqa: BLE001 - the control plane is told the class, never the message
        result = {**base, "outcome": "refused", "reason": "failed", "errorClass": type(error).__name__}
    finally:
        if args.deadline > 0 and hasattr(signal, "SIGALRM"):
            signal.alarm(0)
    (output_dir / "result.json").write_text(json.dumps(result, ensure_ascii=False, sort_keys=True), encoding="utf-8")
    return 0


if __name__ == "__main__":
    sys.exit(main())
