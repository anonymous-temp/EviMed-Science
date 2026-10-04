"""Deterministic checks of a researcher's tables against the meaning recorded for them.

The model states what a dataset means; this module says what the files do. Every
function here is arithmetic over the bytes of a table -- no network, no model, no
clock but the one a caller passes -- so the same two deliveries give the same
findings on any machine, and a finding can be re-derived by anyone who has the
files.

What it answers, each as a named outcome (`DATA_CHECK_OUTCOMES` in
`packages/domain/src/dataSemantics.mjs`; `test_data_semantics_checks.py` holds
this copy equal to the domain's):

- **drift**: the file is not the recorded version; a table, a column appeared or
  went; a column was renamed (a candidate, never applied); its type, its header
  unit, its codes, its scale (a possible unit change), its distribution or its
  missingness moved beyond a stated bound;
- **duplicates**: one observation, by the key the dataset declares, on more than
  one row -- identical rows apart from rows that disagree;
- **joins**: the multiplicity a join was declared with against the one the keys
  actually have, the keys with no partner, a key column that is missing, empty or
  typed differently on the two sides;
- **denominators**: a step that grows when it can only shrink, the size of every
  exclusion, and a step whose count is not the one recorded last time;
- **leakage**: a predictor measured after the time the outcome window opens;
- **transformations**: recorded code that is gone or is not the code that was run.

A finding is information with a name, a count and the rows concerned. It never
stops the rest of an analysis, and a check that could not run says `not_checked`
with the reason -- it never reads as clean.

Privacy: a file's rows stay where they are. What leaves this module is counts,
column names, row numbers and aggregates over at least `MIN_CELL` values; the key
of a duplicated observation is shown only as a pseudonym derived from the file's
own hash, and the control plane's copy of a report drops even that.
"""

from __future__ import annotations

import csv
import hashlib
import io
import math
import os
import re
import stat
from datetime import datetime
from pathlib import PurePosixPath

# --- the vocabulary, as the domain defines it -------------------------------

OUTCOMES = {
    "source_unchanged": ("drift", "information"),
    "source_changed": ("drift", "information"),
    "table_added": ("drift", "attention"),
    "table_removed": ("drift", "attention"),
    "column_added": ("drift", "information"),
    "column_removed": ("drift", "attention"),
    "column_renamed_known": ("drift", "information"),
    "column_renamed_candidate": ("drift", "attention"),
    "type_changed": ("drift", "attention"),
    "unit_changed": ("drift", "attention"),
    "possible_unit_change": ("drift", "attention"),
    "new_codes": ("drift", "attention"),
    "codes_not_seen": ("drift", "information"),
    "distribution_shift": ("drift", "information"),
    "missingness_shift": ("drift", "information"),
    "undeclared_missing_tokens": ("drift", "attention"),
    "duplicate_exact": ("duplicates", "attention"),
    "duplicate_conflicting": ("duplicates", "attention"),
    "join_cardinality_violation": ("joins", "attention"),
    "join_orphans": ("joins", "attention"),
    "join_key_invalid": ("joins", "attention"),
    "denominator_increase": ("denominators", "attention"),
    "denominator_decrease": ("denominators", "information"),
    "denominator_changed": ("denominators", "attention"),
    "temporal_leakage": ("leakage", "attention"),
    "transformation_code_changed": ("transformations", "attention"),
    "transformation_code_missing": ("transformations", "attention"),
}
FAMILIES = ("drift", "duplicates", "joins", "denominators", "leakage", "transformations")
NOT_CHECKED_REASONS = (
    "no_asset", "file_unreadable", "file_too_large", "format_unsupported", "table_unmatched", "no_baseline",
    "observation_key_undeclared", "key_column_missing", "join_table_unavailable", "measurement_time_undeclared",
    "cutoff_undeclared", "time_unparseable", "too_few_values",
)
# `DATA_DRIFT_BOUNDS` in the domain.
MIN_CELL = 10
MEDIAN_SHIFT_IQR = 1.0
SCALE_RATIO = 2.0
MISSING_RATE_DELTA = 0.1
VOCABULARY_MAX = 30
SAMPLE_ROWS = 20
COLUMNS_PROFILED = 300
TYPES = ("integer", "number", "date", "text")

# --- reading ----------------------------------------------------------------

MAX_BYTES = 64 * 1024 * 1024
MAX_ROWS = 2_000_000


class TableError(Exception):
    """A file that cannot be read as tables; `reason` is one of `NOT_CHECKED_REASONS`."""

    def __init__(self, reason: str, message: str):
        super().__init__(message)
        self.reason = reason


class Table:
    """One table of a file: its header, its data rows (with their 1-based numbers) and the file's identity."""

    __slots__ = ("name", "path", "sha256", "bytes", "columns", "rows", "numbers")

    def __init__(self, name, path, sha256, size, columns, rows, numbers):
        self.name = name
        self.path = path
        self.sha256 = sha256
        self.bytes = size
        self.columns = columns
        self.rows = rows
        self.numbers = numbers

    def index(self, column: str):
        try:
            return self.columns.index(column)
        except ValueError:
            return None


def read_workspace_bytes(workspace: str, relative: str, limit=None) -> bytes:
    """One file under the workspace, opened without following a link on the way."""
    limit = MAX_BYTES if limit is None else limit
    parts = PurePosixPath(relative).parts
    if not parts or relative.startswith("/") or "\\" in relative or any(part in ("", ".", "..") for part in parts):
        raise TableError("file_unreadable", "%s is not a workspace-relative path." % relative)
    flags = os.O_RDONLY | getattr(os, "O_NOFOLLOW", 0)
    try:
        directory = os.open(workspace, flags | getattr(os, "O_DIRECTORY", 0))
    except OSError as error:
        raise TableError("file_unreadable", "The workspace is unavailable.") from error
    try:
        for component in parts[:-1]:
            child = os.open(component, flags | getattr(os, "O_DIRECTORY", 0), dir_fd=directory)
            os.close(directory)
            directory = child
        descriptor = os.open(parts[-1], flags | getattr(os, "O_NONBLOCK", 0), dir_fd=directory)
    except OSError as error:
        raise TableError("file_unreadable", "%s was not found in the workspace." % relative) from error
    finally:
        os.close(directory)
    try:
        info = os.fstat(descriptor)
        if not stat.S_ISREG(info.st_mode):
            raise TableError("file_unreadable", "%s is not a regular file." % relative)
        if info.st_size > limit:
            raise TableError("file_too_large", "%s is %d MiB; at most %d MiB can be checked." % (relative, info.st_size // (1024 * 1024) or 1, max(limit // (1024 * 1024), 1)))
        with os.fdopen(descriptor, "rb", closefd=False) as handle:
            return handle.read(limit + 1)
    finally:
        os.close(descriptor)


def _decode(payload: bytes) -> str:
    """UTF-8 (with or without a BOM), else GB18030: Excel on a Chinese machine writes GBK."""
    try:
        return payload.decode("utf-8-sig")
    except UnicodeDecodeError:
        return payload.decode("gb18030", errors="replace")


def _unique_header(header):
    seen = {}
    names = []
    for raw in header:
        name = str(raw).strip()
        count = seen.get(name, 0) + 1
        seen[name] = count
        names.append(name if count == 1 else "%s__%d" % (name, count))
    return names


def _table_from_records(name, path, digest, size, records) -> Table:
    if not records:
        return Table(name, path, digest, size, [], [], [])
    header = _unique_header(records[0])
    width = len(header)
    rows, numbers = [], []
    for position, record in enumerate(records[1:], start=1):
        if not any(str(cell).strip() for cell in record):
            continue
        cells = [str(cell) for cell in record[:width]]
        if len(cells) < width:
            cells.extend([""] * (width - len(cells)))
        rows.append(cells)
        numbers.append(position)
        if len(rows) > MAX_ROWS:
            raise TableError("file_too_large", "%s has more than %d rows." % (path, MAX_ROWS))
    return Table(name, path, digest, size, header, rows, numbers)


def read_tables(workspace: str, relative: str) -> list:
    """The tables of one file: a CSV/TSV is one, a workbook is one per sheet (`file#sheet`).

    Row numbers are 1-based over the records after the header, blank records
    counted, so row N is the (N+1)-th line of a plain CSV and the (N+1)-th row of
    a sheet.
    """
    payload = read_workspace_bytes(workspace, relative)
    digest = hashlib.sha256(payload).hexdigest()
    base = PurePosixPath(relative).name
    suffix = PurePosixPath(relative).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise TableError("format_unsupported", "%s is a workbook and no Excel reader is installed here; export its sheets to CSV." % base) from error
        try:
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        except Exception as error:  # noqa: BLE001 - a damaged workbook is the file's finding, not a crash
            raise TableError("file_unreadable", "%s could not be opened as a workbook." % base) from error
        try:
            tables = []
            for sheet in workbook.worksheets:
                records = [["" if cell is None else str(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
                tables.append(_table_from_records("%s#%s" % (base, sheet.title), relative, digest, len(payload), records))
            return tables
        finally:
            workbook.close()
    if suffix not in (".csv", ".tsv", ".tab", ".txt"):
        raise TableError("format_unsupported", "%s is not a CSV, TSV or Excel file." % base)
    delimiter = "\t" if suffix in (".tsv", ".tab") else ","
    csv.field_size_limit(10 * 1024 * 1024)
    try:
        records = list(csv.reader(io.StringIO(_decode(payload), newline=""), delimiter=delimiter))
    except csv.Error as error:
        raise TableError("file_unreadable", "%s could not be parsed as delimited text." % base) from error
    return [_table_from_records(base, relative, digest, len(payload), records)]


# --- values -----------------------------------------------------------------

INTEGER = re.compile(r"^[+-]?\d+$")
NUMBER = re.compile(r"^[+-]?(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?$")
DATE = re.compile(r"^(?:\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}[-/]\d{1,2}[-/]\d{4})(?:[ tT]\d{1,2}:\d{2}(?::\d{2})?)?\s*$")
# What a spreadsheet or a statistics package writes for "no value" without saying so in a dictionary.
STANDARD_MISSING = frozenset({"na", "n/a", "nan", "null", "none", "-", "--", ".", "?"})


def is_blank(value: str) -> bool:
    return not value or not value.strip()


def is_missing(value: str, tokens) -> bool:
    text = value.strip() if value else ""
    return not text or text in tokens


def is_standard_missing(value: str) -> bool:
    return value.strip().lower() in STANDARD_MISSING


def infer_type(values) -> str:
    """The narrowest type every value satisfies; `empty` for none (the profiler's `infer_type`)."""
    if not values:
        return "empty"
    if all(INTEGER.match(v) for v in values):
        return "integer"
    if all(NUMBER.match(v) for v in values):
        return "number"
    if all(DATE.match(v) for v in values):
        return "date"
    return "text"


def percentile(sorted_values, fraction: float) -> float:
    """Linear interpolation between order statistics (numpy's default)."""
    if len(sorted_values) == 1:
        return sorted_values[0]
    position = (len(sorted_values) - 1) * fraction
    low = math.floor(position)
    high = math.ceil(position)
    return sorted_values[low] + (sorted_values[high] - sorted_values[low]) * (position - low)


def _round(value: float) -> float:
    return float("%.6g" % value)


# A column that names a person or a visit holds identities, whatever else it holds.
IDENTIFIER_TOKENS = frozenset({
    "patient", "subject", "person", "case", "record", "rec", "admission", "visit", "encounter", "inpatient", "outpatient", "mrn",
    "name", "phone", "tel", "telephone", "mobile", "cell", "address", "addr", "email", "mail", "dob", "birth", "birthday",
    "idcard", "ssn", "passport",
})
IDENTIFIER_SUFFIXES = frozenset({"id", "no", "num", "number", "sn", "code"})
IDENTIFIER_WHOLE = frozenset({"id", "mrn", "name", "phone", "tel", "mobile", "address", "email", "dob", "birth", "ssn", "passport", "idcard"})
IDENTIFIER_CJK = re.compile(r"(?:病案|住院|门诊|患者|病人|就诊|身份证|姓名|电话|手机|住址|地址|出生|邮箱|证件)")
MOBILE = re.compile(r"^(?:\+?86)?1[3-9]\d{9}$")
EMAIL = re.compile(r"^[^@\s]+@[^@\s.]+(?:\.[^@\s.]+)+$")
ID_NUMBER = re.compile(r"^\d{17}[\dXx]$")


def name_is_identifying(name: str) -> bool:
    cleaned = name.strip()
    if IDENTIFIER_CJK.search(cleaned):
        return True
    tokens = [t.lower() for t in re.split(r"[^0-9A-Za-z]+", re.sub(r"([a-z])([A-Z])", r"\1_\2", cleaned)) if t]
    if not tokens:
        return False
    if len(tokens) == 1 and tokens[0] in IDENTIFIER_WHOLE:
        return True
    if tokens[-1] in IDENTIFIER_SUFFIXES and any(t in IDENTIFIER_TOKENS for t in tokens[:-1]):
        return True
    return tokens[-1] in IDENTIFIER_WHOLE and len(tokens) > 1 and tokens[-1] != "no"


def values_are_identifying(values) -> bool:
    """A shape that holds for most of a column: an identity number, a mobile number, an e-mail address."""
    sample = [v.strip() for v in values if v and v.strip()]
    if len(sample) < 5:
        return False
    for pattern in (MOBILE, EMAIL, ID_NUMBER):
        if sum(1 for v in sample if pattern.match(v)) / len(sample) >= 0.8:
            return True
    return False


def column_is_identifying(name: str, values, declared_identifier: bool = False) -> bool:
    return declared_identifier or name_is_identifying(name) or values_are_identifying(values)


# A unit written in the header: `Creatinine (umol/L)`, `体重 [kg]`.
HEADER_UNIT = re.compile(r"^(.*?)\s*[(\[（【]\s*([^()\[\]（）【】]{1,32}?)\s*[)\]）】]\s*$")


def split_header(name: str):
    """(base name, unit or None) of a header."""
    match = HEADER_UNIT.match(name.strip())
    if not match or not match.group(1).strip():
        return name.strip(), None
    return match.group(1).strip(), match.group(2).strip()


def unit_key(unit) -> str:
    """A unit compared by what it says: case, spacing and the two micro signs do not matter."""
    if unit is None:
        return ""
    return re.sub(r"\s+", "", str(unit)).lower().replace("µ", "u").replace("μ", "u").rstrip(".")


# Common clinical units that differ by a fixed factor: `value in a ≈ value in b × factor`. A scale shift
# that matches one of these is named as a possible unit change; any other shift stays a distribution shift.
UNIT_FACTORS = (
    ("umol/L", "mg/dL", 88.4, "creatinine"),
    ("mg/dL", "mmol/L", 18.016, "glucose"),
    ("mg/dL", "mmol/L", 38.67, "total, LDL or HDL cholesterol"),
    ("mg/dL", "mmol/L", 88.57, "triglycerides"),
    ("mg/dL", "mmol/L", 2.801, "urea nitrogen"),
    ("umol/L", "mg/dL", 17.1, "bilirubin"),
    ("umol/L", "mg/dL", 59.48, "uric acid"),
    ("mg/dL", "mmol/L", 4.008, "calcium"),
    ("g/L", "g/dL", 10.0, "hemoglobin or albumin"),
    ("lb", "kg", 2.2046, "weight"),
    ("cm", "in", 2.54, "height or length"),
)
FACTOR_TOLERANCE = 0.03


def unit_candidates(ratio: float, declared_unit):
    """Unit pairs whose factor this median ratio matches, the declared unit's first.

    `ratio` is new/old. A ratio near the factor means the values were multiplied by it, so the old
    unit was the one with the smaller numbers.
    """
    found = []
    for big, small, factor, what in UNIT_FACTORS:
        for old, new, expected in ((small, big, factor), (big, small, 1.0 / factor)):
            if abs(ratio - expected) / expected <= FACTOR_TOLERANCE:
                hint = "declared_unit_matches" if declared_unit and unit_key(declared_unit) == unit_key(old) else "by_factor_only"
                found.append({"from": old, "to": new, "factor": factor, "measure": what, "hint": hint})
    found.sort(key=lambda item: (item["hint"] != "declared_unit_matches", item["measure"]))
    return found


def parse_time(value: str):
    """(datetime, has_time) from the closed set of unambiguous forms, else None.

    Day-first and month-first both read `03/04/2024`, so neither is read: a date
    that cannot say which it is stays unparsed and is counted, not guessed.
    """
    text = value.strip().replace("T", " ")
    for pattern, has_time in (
        ("%Y-%m-%d %H:%M:%S", True), ("%Y-%m-%d %H:%M", True), ("%Y/%m/%d %H:%M:%S", True), ("%Y/%m/%d %H:%M", True),
        ("%Y-%m-%d", False), ("%Y/%m/%d", False), ("%Y%m%d", False),
    ):
        try:
            return datetime.strptime(text, pattern), has_time
        except ValueError:
            continue
    return None


# --- the profile of one table (what a binding records) -----------------------

def profile_column(name, cells, missing_tokens=(), identifying=False) -> dict:
    """Aggregates of one column: its observed type, how many are missing, how many distinct, a numeric
    summary when at least `MIN_CELL` numbers stand behind it, and the unit its header states. Never a value."""
    tokens = set(missing_tokens)
    filled = [c.strip() for c in cells if not is_missing(c, tokens) and not is_standard_missing(c)]
    kind = infer_type(filled)
    profile = {"name": name, "type": kind, "missing": len(cells) - len(filled), "distinct": len(set(filled))}
    _base, unit = split_header(name)
    if unit:
        profile["headerUnit"] = unit
    if kind in ("integer", "number") and len(filled) >= MIN_CELL and not identifying:
        numbers = sorted(float(v) for v in filled)
        profile["numeric"] = {
            "n": len(numbers), "min": _round(numbers[0]), "p25": _round(percentile(numbers, 0.25)), "median": _round(percentile(numbers, 0.5)),
            "p75": _round(percentile(numbers, 0.75)), "max": _round(numbers[-1]), "mean": _round(sum(numbers) / len(numbers)),
        }
    return profile


def binding_of(table: Table, declared: dict | None = None) -> dict:
    """What a binding records of one table: the file's hash and size, the row count and the column profiles."""
    declared = declared or {}
    identifier_columns = declared.get("identifierColumns", set())
    columns = []
    for index, name in enumerate(table.columns[:COLUMNS_PROFILED]):
        cells = [row[index] for row in table.rows]
        variable = declared.get("variables", {}).get(name, {})
        identifying = column_is_identifying(name, cells, name in identifier_columns or variable.get("role") == "identifier")
        columns.append(profile_column(name, cells, variable.get("missingTokens", ()), identifying))
    return {"table": table.name, "path": table.path, "sha256": table.sha256, "bytes": table.bytes, "rows": len(table.rows), "columns": columns}


def observed_vocabulary(table: Table, column: str, declared: dict | None = None):
    """The distinct non-missing values of a coded column, or why they are not given.

    An identifier is never listed, and neither is a column too varied to be a code list: the value
    a model writes down as `allowedValues` is copied from here, not typed.
    """
    index = table.index(column)
    if index is None:
        return None, "key_column_missing"
    declared = declared or {}
    variable = declared.get("variables", {}).get(column, {})
    cells = [row[index] for row in table.rows]
    if column_is_identifying(column, cells, column in declared.get("identifierColumns", set()) or variable.get("role") == "identifier"):
        return None, "identifying"
    tokens = set(variable.get("missingTokens", ()))
    values = sorted({c.strip() for c in cells if not is_missing(c, tokens) and not is_standard_missing(c)})
    if len(values) > VOCABULARY_MAX or any(len(v) > 40 for v in values):
        return None, "high_cardinality"
    return values, None


# --- the recorded meaning, flattened for the checks --------------------------

def resolve(asset: dict | None, overrides: dict | None = None) -> dict:
    """The asset as plain values (no basis, no history), overlaid with what the caller declared inline.

    {tables: {name: {subjectKey, observationKey, variables: {col: {...}}, identifierColumns}},
     joins: [...], bindings: {table: binding}, aliases: {table: {current: [former...]}},
     denominators: {...}, transformations: [...]}
    """
    overrides = overrides or {}
    resolved = {"tables": {}, "joins": [], "bindings": {}, "denominators": {}, "transformations": []}
    if asset:
        for table in asset.get("tables", []):
            entry = {"subjectKey": None, "observationKey": None, "variables": {}, "identifierColumns": set()}
            for facet in ("subjectKey", "observationKey"):
                fact = table.get("facts", {}).get(facet)
                if fact:
                    entry[facet] = list(fact["value"])
            for variable in table.get("variables", []):
                facts = variable.get("facts", {})
                record = {key: facts[key]["value"] for key in ("type", "unit", "role", "valueSource", "range", "aliases", "allowedValues") if key in facts}
                if "allowedValues" in record:
                    record["allowedCodes"] = [item["code"] for item in record["allowedValues"]]
                if "missingness" in facts:
                    record["missingTokens"] = list(facts["missingness"]["value"].get("tokens", []))
                if "measuredAt" in facts:
                    record["measuredAt"] = facts["measuredAt"]["value"]["column"]
                entry["variables"][variable["name"]] = record
                if record.get("role") == "identifier":
                    entry["identifierColumns"].add(variable["name"])
            resolved["tables"][table["name"]] = entry
        for join in asset.get("joins", []):
            cardinality = join.get("facts", {}).get("cardinality")
            resolved["joins"].append({"id": join["id"], "left": join["left"], "right": join["right"], "cardinality": cardinality["value"] if cardinality else None})
        resolved["bindings"] = {binding["table"]: binding for binding in asset.get("bindings", [])}
        resolved["denominators"] = dict(asset.get("denominators", {}))
        resolved["transformations"] = list(asset.get("transformations", []))
    for kind in ("observationKey", "subjectKey"):
        for table, columns in (overrides.get(kind + "s") or {}).items():
            resolved["tables"].setdefault(table, {"subjectKey": None, "observationKey": None, "variables": {}, "identifierColumns": set()})[kind] = list(columns)
    for join in overrides.get("joins") or []:
        resolved["joins"].append({"id": "%s.%s->%s.%s" % (join["left"]["table"], "+".join(join["left"]["columns"]), join["right"]["table"], "+".join(join["right"]["columns"])),
                                  "left": join["left"], "right": join["right"], "cardinality": join.get("cardinality")})
    for table in resolved["tables"].values():
        table["identifierColumns"] |= {column for key in ("subjectKey", "observationKey") for column in (table[key] or [])}
    return resolved


def _former_names(declared_table: dict) -> dict:
    """{former column name: current name} from the variables' recorded aliases."""
    return {alias: name for name, record in declared_table.get("variables", {}).items() for alias in record.get("aliases", [])}


def resolve_columns(table: Table, wanted, declared_table: dict):
    """The file's columns for declared names: the name itself, or the current name of a variable that used to carry it."""
    former = _former_names(declared_table)
    found, missing = [], []
    for name in wanted:
        if name in table.columns:
            found.append(name)
        elif former.get(name) in table.columns:
            found.append(former[name])
        else:
            missing.append(name)
    return found, missing


# --- findings ---------------------------------------------------------------

def finding(outcome: str, subject: dict, message: str, count=None, rows=None, detail=None, examples=None, severity=None) -> dict:
    family, default = OUTCOMES[outcome]
    item = {"outcome": outcome, "family": family, "severity": severity or default, "subject": {k: v for k, v in subject.items() if v}, "message": message[:300]}
    if count is not None:
        item["count"] = int(count)
    if rows:
        item["rows"] = [int(r) for r in rows[:SAMPLE_ROWS]]
    if detail:
        item["detail"] = detail
    if examples:
        item["examples"] = examples
    return item


def _not_checked(family: str, reason: str, subject: dict, message: str) -> dict:
    return {"family": family, "reason": reason, "subject": {k: v for k, v in subject.items() if v}, "message": message}


def _clean(family: str, subject: dict) -> dict:
    return {"family": family, "subject": {k: v for k, v in subject.items() if v}}


def _ref(salt: str, key) -> str:
    """A short pseudonym of a key, derived from the file's own hash: stable within the file, meaningless outside it."""
    return hashlib.sha256((salt + "\0" + "\0".join(key)).encode("utf-8")).hexdigest()[:8]


# --- drift ------------------------------------------------------------------

def _cells(table: Table, column: str):
    index = table.index(column)
    return [] if index is None else [row[index] for row in table.rows]


def check_drift(table: Table, previous: dict | None, declared_table: dict, found: list, notes: list, clean: list) -> None:
    """The new delivery of one table against the recorded version and the recorded meaning."""
    subject = {"table": table.name}
    if previous is None:
        notes.append(_not_checked("drift", "no_baseline", subject, "No earlier version of %s is recorded, so there is nothing to compare its columns with." % table.name))
        previous_columns = {}
    else:
        if previous["sha256"] == table.sha256:
            found.append(finding("source_unchanged", subject, "%s is byte-identical to the recorded version (%d rows)." % (table.name, len(table.rows)), detail={"rows": len(table.rows)}))
            return
        found.append(finding("source_changed", subject, "%s is not the recorded version: %d rows now, %d before." % (table.name, len(table.rows), previous["rows"]),
                             detail={"rowsBefore": previous["rows"], "rowsAfter": len(table.rows), "bytesBefore": previous["bytes"], "bytesAfter": table.bytes}))
        previous_columns = {column["name"]: column for column in previous.get("columns", [])}
    variables = declared_table.get("variables", {})
    former = _former_names(declared_table)
    new_profiles = {column["name"]: column for column in binding_of(table, declared_table)["columns"]}
    new_names = set(table.columns)

    removed = [name for name in previous_columns if name not in new_names]
    added = [name for name in table.columns if name not in previous_columns]
    paired = {}  # new name -> previous name, for columns that carry on under another header
    # A header whose unit changed is the same column, not a removed one and an added one.
    for old in list(removed):
        base, unit = split_header(old)
        for new in list(added):
            new_base, new_unit = split_header(new)
            if new_base == base and (unit or new_unit) and unit_key(unit) != unit_key(new_unit):
                paired[new] = old
                removed.remove(old)
                added.remove(new)
                found.append(finding("unit_changed", {"table": table.name, "column": base},
                                     "The header of %s changed its unit from %s to %s." % (base, unit or "none stated", new_unit or "none stated"),
                                     detail={"from": unit, "to": new_unit, "evidence": "header"}))
                break
    # A name the data used to go by is a rename the dataset already knows about.
    for old in list(removed):
        for new in list(added):
            if old in variables.get(new, {}).get("aliases", []) or former.get(old) == new:
                paired[new] = old
                removed.remove(old)
                added.remove(new)
                found.append(finding("column_renamed_known", {"table": table.name, "column": new},
                                     "%s is recorded as the former name of %s." % (old, new), detail={"from": old, "to": new}))
                break
    # What is left: a candidate rename is a removed and an added column that look like one column.
    old_position = {column["name"]: index for index, column in enumerate(previous["columns"])} if previous else {}
    candidates = []
    for old in removed:
        for new in added:
            old_profile, new_profile = previous_columns[old], new_profiles[new]
            if old_profile["type"] != new_profile["type"] or old_profile["type"] == "empty":
                continue
            evidence = ["same type"]
            score = 1
            if old_position.get(old) == table.columns.index(new):
                evidence.append("same position")
                score += 2
            a, b = old_profile.get("numeric"), new_profile.get("numeric")
            if a and b and a["median"] and abs(b["median"] - a["median"]) <= 0.1 * abs(a["median"]):
                evidence.append("similar median")
                score += 1
            if score >= 3:
                candidates.append((score, old, new, evidence))
    candidates.sort(key=lambda item: (-item[0], item[1], item[2]))
    for _score, old, new, evidence in candidates:
        if old not in removed or new not in added:
            continue
        removed.remove(old)
        added.remove(new)
        paired[new] = old
        found.append(finding("column_renamed_candidate", {"table": table.name, "column": new},
                             "%s is missing and %s is new; they look like one column renamed (%s). Confirm before relying on it." % (old, new, ", ".join(evidence)),
                             detail={"from": old, "to": new, "evidence": evidence}))
    for name in removed:
        found.append(finding("column_removed", {"table": table.name, "column": name}, "The recorded column %s is not in this delivery." % name))
    if added:
        found.append(finding("column_added", {"table": table.name}, "%d new column(s): %s." % (len(added), ", ".join(added[:10])),
                             count=len(added), detail={"columns": added[:20]}))

    # Columns that carry on, under the same name or a recognised one.
    for name in table.columns:
        old_name = paired.get(name, name if name in previous_columns else None)
        record = variables.get(name) or variables.get(old_name or "", {})
        profile = new_profiles[name]
        old = previous_columns.get(old_name) if old_name else None
        subject = {"table": table.name, "column": name}
        cells = _cells(table, name)
        tokens = set(record.get("missingTokens", ()))
        identifying = column_is_identifying(name, cells, name in declared_table.get("identifierColumns", set()) or record.get("role") == "identifier")
        if old and old["type"] != "empty" and profile["type"] != "empty" and old["type"] != profile["type"]:
            found.append(finding("type_changed", subject, "%s was %s and now reads as %s." % (name, old["type"], profile["type"]),
                                 detail={"from": old["type"], "to": profile["type"]}))
        elif not old and record.get("type") and profile["type"] not in ("empty", record["type"]):
            found.append(finding("type_changed", subject, "%s is recorded as %s and now reads as %s." % (name, record["type"], profile["type"]),
                                 detail={"from": record["type"], "to": profile["type"], "against": "recorded type"}))
        # The header's own statement of the unit, against the recorded one.
        _base, unit = split_header(name)
        if unit and record.get("unit") and unit_key(unit) != unit_key(record["unit"]) and name not in paired:
            found.append(finding("unit_changed", subject, "The header of %s says %s; the recorded unit is %s." % (name, unit, record["unit"]),
                                 detail={"from": record["unit"], "to": unit, "evidence": "header"}))
        # A tokens column is numeric and carries a marker nobody declared.
        if profile["type"] in ("integer", "number") and not identifying:
            index = table.index(name)
            stray = {}
            first_rows = []
            for row, number in zip(table.rows, table.numbers):
                cell = row[index].strip()
                if cell and cell not in tokens and is_standard_missing(cell):
                    stray[cell] = stray.get(cell, 0) + 1
                    if len(first_rows) < SAMPLE_ROWS:
                        first_rows.append(number)
            if stray:
                found.append(finding("undeclared_missing_tokens", subject,
                                     "%s is numeric but %d value(s) are %s, which no missingness declaration covers." % (name, sum(stray.values()), ", ".join(sorted(stray))),
                                     count=sum(stray.values()), rows=first_rows, detail={"tokens": sorted(stray)[:5]}))
        if old and old.get("numeric") and profile.get("numeric") and not identifying:
            a, b = old["numeric"], profile["numeric"]
            ratio = b["median"] / a["median"] if a["median"] and b["median"] and (a["median"] > 0) == (b["median"] > 0) else None
            iqr = a["p75"] - a["p25"]
            shift = abs(b["median"] - a["median"]) / iqr if iqr > 0 else None
            if ratio is not None and (ratio >= SCALE_RATIO or ratio <= 1 / SCALE_RATIO):
                matches = unit_candidates(ratio, record.get("unit"))
                if matches:
                    best = matches[0]
                    found.append(finding("possible_unit_change", subject,
                                         "Values of %s are %.3g times what they were (median %s -> %s); that is the factor between %s and %s (%s)." % (
                                             name, ratio, a["median"], b["median"], best["from"], best["to"], best["measure"]),
                                         detail={"ratio": _round(ratio), "medianBefore": a["median"], "medianAfter": b["median"], "from": best["from"], "to": best["to"],
                                                 "measure": best["measure"], "hint": best["hint"]}))
                else:
                    found.append(finding("distribution_shift", subject,
                                         "Values of %s are %.3g times what they were (median %s -> %s), beyond the %gx bound; no known unit conversion has that factor." % (
                                             name, ratio, a["median"], b["median"], SCALE_RATIO),
                                         detail={"ratio": _round(ratio), "medianBefore": a["median"], "medianAfter": b["median"], "bound": "scaleRatio"}))
            elif shift is not None and shift > MEDIAN_SHIFT_IQR:
                found.append(finding("distribution_shift", subject,
                                     "The median of %s moved from %s to %s, %.2g interquartile ranges of the earlier delivery." % (name, a["median"], b["median"], shift),
                                     detail={"medianBefore": a["median"], "medianAfter": b["median"], "iqrShifts": _round(shift), "bound": "medianShiftIqr"}))
        if old and previous and previous["rows"] and table.rows:
            before = old["missing"] / previous["rows"]
            after = profile["missing"] / len(table.rows)
            if abs(after - before) > MISSING_RATE_DELTA:
                found.append(finding("missingness_shift", subject,
                                     "The share of missing values in %s moved from %.0f%% to %.0f%%." % (name, before * 100, after * 100),
                                     detail={"before": _round(before), "after": _round(after), "bound": "missingRateDelta"}))
        allowed = record.get("allowedCodes")
        if allowed and not identifying:
            index = table.index(name)
            observed = {}
            sample = {}
            for row, number in zip(table.rows, table.numbers):
                cell = row[index].strip()
                if is_missing(cell, tokens) or is_standard_missing(cell):
                    continue
                if cell not in allowed:
                    observed[cell] = observed.get(cell, 0) + 1
                    sample.setdefault(cell, number)
            if observed:
                top = sorted(observed.items(), key=lambda item: (-item[1], item[0]))[:20]
                found.append(finding("new_codes", subject,
                                     "%s has %d code(s) that are not in its recorded list: %s." % (name, len(observed), ", ".join(code for code, _ in top[:8])),
                                     count=sum(observed.values()), rows=sorted(sample.values()), detail={"codes": [code for code, _ in top], "distinctNew": len(observed)}))
            seen = {row[index].strip() for row in table.rows}
            unseen = [code for code in allowed if code not in seen]
            if unseen:
                found.append(finding("codes_not_seen", subject, "%d recorded code(s) of %s do not occur in this delivery: %s." % (len(unseen), name, ", ".join(unseen[:8])),
                                     count=len(unseen), detail={"codes": unseen[:20]}))
    if not any(item["family"] == "drift" and item["outcome"] not in ("source_changed",) and item["subject"].get("table") == table.name for item in found):
        clean.append(_clean("drift", {"table": table.name}))


# --- duplicates -------------------------------------------------------------

def check_duplicates(table: Table, declared_table: dict, found: list, notes: list, clean: list, warnings: list) -> None:
    subject = {"table": table.name}
    wanted = declared_table.get("observationKey")
    if not wanted:
        notes.append(_not_checked("duplicates", "observation_key_undeclared", subject,
                                  "%s has no declared observation key, so repeated observations were not looked for. Declare which columns identify one observation." % table.name))
        return
    columns, missing = resolve_columns(table, wanted, declared_table)
    if missing:
        notes.append(_not_checked("duplicates", "key_column_missing", {**subject, "column": ",".join(missing)},
                                  "The observation key column(s) %s are not in %s (renamed? record the former name as an alias)." % (", ".join(missing), table.name)))
        return
    key_indexes = [table.columns.index(c) for c in columns]
    groups = {}
    blank_keys = 0
    for position, row in enumerate(table.rows):
        key = tuple(row[i].strip() for i in key_indexes)
        if any(not part for part in key):
            blank_keys += 1
            continue
        groups.setdefault(key, []).append(position)
    exact_rows, exact_groups, conflict_rows, conflict_groups = [], [], [], []
    differing = {}
    for key, positions in groups.items():
        if len(positions) < 2:
            continue
        normalised = [tuple(cell.strip() for cell in table.rows[p]) for p in positions]
        if len(set(normalised)) == 1:
            exact_groups.append((key, positions))
            exact_rows.extend(positions[1:])
        else:
            conflict_groups.append((key, positions))
            conflict_rows.extend(positions[1:])
            for index, name in enumerate(table.columns):
                if len({row[index] for row in normalised}) > 1:
                    differing[name] = differing.get(name, 0) + 1
    label = "+".join(columns)

    def examples(group_list):
        return [{"key": _ref(table.sha256, key), "rows": [table.numbers[p] for p in positions[:6]]} for key, positions in group_list[:5]]

    if exact_groups:
        found.append(finding("duplicate_exact", {**subject, "column": label},
                             "%d observation(s) by %s appear on more than one identical row (%d extra row(s)); nothing was removed." % (len(exact_groups), label, len(exact_rows)),
                             count=len(exact_rows), rows=[table.numbers[p] for p in exact_rows],
                             detail={"groups": len(exact_groups), "extraRows": len(exact_rows)}, examples=examples(exact_groups)))
    if conflict_groups:
        columns_that_differ = [name for name, _count in sorted(differing.items(), key=lambda item: (-item[1], item[0]))][:10]
        found.append(finding("duplicate_conflicting", {**subject, "column": label},
                             "%d observation(s) by %s appear on several rows that disagree (in %s); none was chosen over another." % (len(conflict_groups), label, ", ".join(columns_that_differ[:5])),
                             count=len(conflict_rows), rows=[table.numbers[p] for p in conflict_rows],
                             detail={"groups": len(conflict_groups), "extraRows": len(conflict_rows), "differingColumns": columns_that_differ}, examples=examples(conflict_groups)))
    if not exact_groups and not conflict_groups:
        clean.append(_clean("duplicates", {**subject, "column": label}))
    if blank_keys:
        warnings.append("%d row(s) of %s have a blank part in the observation key and were not compared." % (blank_keys, table.name))


# --- joins ------------------------------------------------------------------

WIDEST = {"one_to_one": (1, 1), "one_to_many": (1, None), "many_to_one": (None, 1), "many_to_many": (None, None)}


def _observed_cardinality(left_max: int, right_max: int) -> str:
    if left_max <= 1 and right_max <= 1:
        return "one_to_one"
    if left_max <= 1:
        return "one_to_many"
    if right_max <= 1:
        return "many_to_one"
    return "many_to_many"


def check_join(join: dict, tables: dict, declared: dict, found: list, notes: list, clean: list, warnings: list, profiles: list) -> None:
    left_table, right_table = tables.get(join["left"]["table"]), tables.get(join["right"]["table"])
    subject = {"join": join["id"]}
    if left_table is None or right_table is None:
        absent = join["left"]["table"] if left_table is None else join["right"]["table"]
        notes.append(_not_checked("joins", "join_table_unavailable", subject, "The join needs %s, which was not among the files checked." % absent))
        return
    sides = []
    for side, table in ((join["left"], left_table), (join["right"], right_table)):
        columns, missing = resolve_columns(table, side["columns"], declared.get("tables", {}).get(table.name, {}))
        if missing:
            found.append(finding("join_key_invalid", {**subject, "table": table.name, "column": ",".join(missing)},
                                 "Join key column(s) %s are not in %s." % (", ".join(missing), table.name), detail={"problem": "column_missing"}))
            return
        sides.append((table, columns))
    keys = []
    for table, columns in sides:
        indexes = [table.columns.index(c) for c in columns]
        counter, rows_of, blank = {}, {}, 0
        for position, row in enumerate(table.rows):
            key = tuple(row[i].strip() for i in indexes)
            if any(not part for part in key):
                blank += 1
                continue
            counter[key] = counter.get(key, 0) + 1
            rows_of.setdefault(key, []).append(position)
        keys.append((counter, rows_of, blank, [infer_type([part for part in {row[i].strip() for row in table.rows} if part]) for i in indexes]))
    (left_count, left_rows, left_blank, left_types), (right_count, right_rows, right_blank, right_types) = keys
    for (table, _columns), counter, blank in ((sides[0], left_count, left_blank), (sides[1], right_count, right_blank)):
        if table.rows and not counter:
            found.append(finding("join_key_invalid", {**subject, "table": table.name}, "Every join key in %s is blank." % table.name, detail={"problem": "all_blank", "rows": len(table.rows)}))
            return
    family = {"integer": "numeric", "number": "numeric"}
    if [family.get(t, t) for t in left_types] != [family.get(t, t) for t in right_types] and "empty" not in left_types + right_types:
        found.append(finding("join_key_invalid", subject, "The key is typed %s in %s and %s in %s, so equal-looking keys may not match." % (
            "/".join(left_types), left_table.name, "/".join(right_types), right_table.name), detail={"problem": "type_mismatch", "left": "/".join(left_types), "right": "/".join(right_types)}))
    left_max = max(left_count.values(), default=0)
    right_max = max(right_count.values(), default=0)
    observed = _observed_cardinality(left_max, right_max)
    declared_kind = join.get("cardinality")
    problems = False
    if declared_kind in WIDEST:
        left_unique, right_unique = WIDEST[declared_kind][0] == 1, WIDEST[declared_kind][1] == 1
        offending = []
        if left_unique and left_max > 1:
            offending.append(("left", left_count, left_rows, left_table))
        if right_unique and right_max > 1:
            offending.append(("right", right_count, right_rows, right_table))
        for side, counter, rows_of, table in offending:
            repeated = {key: positions for key, positions in rows_of.items() if len(positions) > 1}
            extra = [p for positions in repeated.values() for p in positions[1:]]
            fan = sum(left_count[key] * right_count[key] for key in left_count if key in right_count)
            found.append(finding("join_cardinality_violation", {**subject, "table": table.name},
                                 "The join is declared %s but %d key(s) repeat on the %s side (%s); joining as declared would give %d rows from %d." % (
                                     declared_kind, len(repeated), side, table.name, fan, sum(left_count.values())),
                                 count=len(extra), rows=[table.numbers[p] for p in extra],
                                 detail={"declared": declared_kind, "observed": observed, "side": side, "repeatedKeys": len(repeated), "rowsAfterJoin": fan, "leftRows": sum(left_count.values())},
                                 examples=[{"key": _ref(table.sha256, key), "rows": [table.numbers[p] for p in positions[:6]]} for key, positions in list(repeated.items())[:5]]))
            problems = True
    left_only = [key for key in left_count if key not in right_count]
    right_only = [key for key in right_count if key not in left_count]
    if left_only or right_only:
        left_orphan_rows = [p for key in left_only for p in left_rows[key]]
        right_orphan_rows = [p for key in right_only for p in right_rows[key]]
        # An unmatched key on the "many" side has lost its parent; on the "one" side it only has no children.
        severe = (left_only and declared_kind in ("many_to_one", "one_to_one", "many_to_many", None)) or (right_only and declared_kind in ("one_to_many", "one_to_one", "many_to_many", None))
        found.append(finding("join_orphans", subject,
                             "%d key(s) in %s have no partner in %s (%d rows); %d key(s) in %s have none in %s (%d rows)." % (
                                 len(left_only), left_table.name, right_table.name, len(left_orphan_rows), len(right_only), right_table.name, left_table.name, len(right_orphan_rows)),
                             count=len(left_orphan_rows) + len(right_orphan_rows),
                             rows=[left_table.numbers[p] for p in left_orphan_rows] if left_orphan_rows else [right_table.numbers[p] for p in right_orphan_rows],
                             detail={"leftOnlyKeys": len(left_only), "leftOnlyRows": len(left_orphan_rows), "rightOnlyKeys": len(right_only), "rightOnlyRows": len(right_orphan_rows),
                                     "leftRows": sum(left_count.values()), "rightRows": sum(right_count.values())},
                             severity="attention" if severe else "information"))
        problems = True
    if left_blank or right_blank:
        warnings.append("%d row(s) of %s and %d of %s have a blank join key and match nothing." % (left_blank, left_table.name, right_blank, right_table.name))
    profiles.append({"join": join["id"], "declared": declared_kind, "observed": observed, "leftKeys": len(left_count), "rightKeys": len(right_count),
                     "leftRows": sum(left_count.values()), "rightRows": sum(right_count.values())})
    if not problems:
        clean.append(_clean("joins", subject))


# --- denominators -----------------------------------------------------------

MONOTONE_ROWS = ("filter", "dedupe", "recode")
MONOTONE_SUBJECTS = ("filter", "dedupe", "recode", "derive")


def check_denominators(steps: list, tables: dict, declared: dict, recorded: dict, found: list, notes: list, clean: list) -> dict:
    """The size of every step, the steps that grew, and the ones that are not what was recorded.

    Returns the denominators observed, by label, for the next check to compare with.
    """
    observed = []
    for step in steps:
        label = str(step["label"])
        rows = step.get("rows")
        subjects = step.get("subjects")
        source = "reported"
        if rows is None:
            table = tables.get(step.get("table") or "")
            if table is None:
                notes.append(_not_checked("denominators", "table_unmatched", {"step": label}, "Step %s names a table that was not read, and reported no counts." % label))
                continue
            rows, source = len(table.rows), "measured"
            wanted = step.get("subjectColumns") or declared.get("tables", {}).get(table.name, {}).get("subjectKey")
            if wanted:
                columns, missing = resolve_columns(table, wanted, declared.get("tables", {}).get(table.name, {}))
                if not missing:
                    indexes = [table.columns.index(c) for c in columns]
                    subjects = len({tuple(row[i].strip() for i in indexes) for row in table.rows if all(row[i].strip() for i in indexes)})
        observed.append({"label": label, "kind": step.get("kind", "filter"), "rows": int(rows), "subjects": None if subjects is None else int(subjects), "source": source})
    for before, after in zip(observed, observed[1:]):
        subject = {"step": "%s -> %s" % (before["label"], after["label"])}
        pairs = (("rows", before["rows"], after["rows"], MONOTONE_ROWS), ("subjects", before["subjects"], after["subjects"], MONOTONE_SUBJECTS))
        grew = [(what, was, now) for what, was, now, monotone in pairs if was is not None and now is not None and now > was and after["kind"] in monotone]
        shrank = [(what, was, now) for what, was, now, _monotone in pairs if was is not None and now is not None and now < was]
        counts = {"rowsBefore": before["rows"], "rowsAfter": after["rows"], "subjectsBefore": before["subjects"], "subjectsAfter": after["subjects"]}
        if grew:
            found.append(finding("denominator_increase", subject,
                                 "%s has more %s than %s (%s), though a %s step can only remove them." % (
                                     after["label"], " and ".join(w for w, _a, _b in grew), before["label"], ", ".join("%d against %d" % (n, w) for _x, w, n in grew), after["kind"]),
                                 count=max(n - w for _x, w, n in grew), detail={**counts, "what": [w for w, _a, _b in grew], "kind": after["kind"]}))
        if shrank:
            what, was, now = shrank[0]
            found.append(finding("denominator_decrease", subject,
                                 "%s: %s." % (
                                     "%s to %s" % (before["label"], after["label"]),
                                     "; ".join("%d of %d %s (%.1f%%) are gone" % (w - n, w, x, (w - n) / w * 100) for x, w, n in shrank)),
                                 count=was - now, detail={**counts, "what": [x for x, _w, _n in shrank], "share": _round((was - now) / was)}))
    for item in observed:
        was = recorded.get(item["label"])
        subject = {"step": item["label"]}
        if was and (was["rows"] != item["rows"] or (was.get("subjects") is not None and item["subjects"] is not None and was["subjects"] != item["subjects"])):
            found.append(finding("denominator_changed", subject,
                                 "%s now has %d rows%s; the last analysis recorded %d%s." % (
                                     item["label"], item["rows"], "" if item["subjects"] is None else " and %d subjects" % item["subjects"], was["rows"],
                                     "" if was.get("subjects") is None else " and %d subjects" % was["subjects"]),
                                 detail={"rowsBefore": was["rows"], "rowsAfter": item["rows"], "subjectsBefore": was.get("subjects"), "subjectsAfter": item["subjects"], "source": item["source"]}))
        elif was:
            clean.append(_clean("denominators", subject))
    return {item["label"]: {"rows": item["rows"], "subjects": item["subjects"], "source": item["source"]} for item in observed}


# --- leakage ----------------------------------------------------------------

def _entity_times(table: Table, key_columns, time_column: str):
    indexes = [table.columns.index(c) for c in key_columns]
    time_index = table.columns.index(time_column)
    earliest, unparseable = {}, 0
    for row in table.rows:
        key = tuple(row[i].strip() for i in indexes)
        if any(not part for part in key) or not row[time_index].strip():
            continue
        parsed = parse_time(row[time_index])
        if parsed is None:
            unparseable += 1
            continue
        current = earliest.get(key)
        if current is None or parsed[0] < current[0]:
            earliest[key] = parsed
    return earliest, unparseable


def check_leakage(request: dict, tables: dict, declared: dict, found: list, notes: list, clean: list) -> None:
    cutoff = request.get("cutoff")
    if not cutoff:
        notes.append(_not_checked("leakage", "cutoff_undeclared", {}, "No cutoff was given: say which column holds the time the outcome window opens (or the outcome time itself)."))
        return
    cutoff_table = tables.get(cutoff.get("table") or "")
    if cutoff_table is None or cutoff["column"] not in cutoff_table.columns:
        notes.append(_not_checked("leakage", "cutoff_undeclared", {"table": cutoff.get("table"), "column": cutoff.get("column")}, "The cutoff column %s is not in the files checked." % cutoff.get("column")))
        return
    subject_columns = request.get("subjectColumns") or declared.get("tables", {}).get(cutoff_table.name, {}).get("subjectKey")
    if not subject_columns:
        notes.append(_not_checked("leakage", "observation_key_undeclared", {"table": cutoff_table.name}, "Leakage is judged per subject and %s declares no subject key." % cutoff_table.name))
        return
    own, missing = resolve_columns(cutoff_table, subject_columns, declared.get("tables", {}).get(cutoff_table.name, {}))
    if missing:
        notes.append(_not_checked("leakage", "key_column_missing", {"table": cutoff_table.name, "column": ",".join(missing)}, "The subject key column(s) %s are not in %s." % (", ".join(missing), cutoff_table.name)))
        return
    cutoffs, bad_cutoffs = _entity_times(cutoff_table, own, cutoff["column"])
    if not cutoffs:
        notes.append(_not_checked("leakage", "time_unparseable", {"table": cutoff_table.name, "column": cutoff["column"]},
                                  "No value of %s reads as a date or time (%d did not)." % (cutoff["column"], bad_cutoffs)))
        return
    predictors = request.get("predictors")
    if not predictors:
        predictors = []
        for name, table in declared.get("tables", {}).items():
            for column, record in table.get("variables", {}).items():
                if record.get("role") in ("exposure", "covariate") and record.get("measuredAt"):
                    predictors.append({"table": name, "column": column})
    if not predictors:
        notes.append(_not_checked("leakage", "measurement_time_undeclared", {}, "No predictor was named and none has a recorded measurement time column."))
        return
    for predictor in predictors:
        table = tables.get(predictor.get("table") or cutoff_table.name)
        column = predictor["column"]
        subject = {"table": table.name if table else predictor.get("table"), "predictor": column}
        if table is None or column not in table.columns:
            notes.append(_not_checked("leakage", "table_unmatched", subject, "Predictor %s is not in the files checked." % column))
            continue
        time_column = predictor.get("timeColumn") or declared.get("tables", {}).get(table.name, {}).get("variables", {}).get(column, {}).get("measuredAt")
        if not time_column or time_column not in table.columns:
            notes.append(_not_checked("leakage", "measurement_time_undeclared", subject, "No column says when %s was measured." % column))
            continue
        keys, missing = resolve_columns(table, subject_columns, declared.get("tables", {}).get(table.name, {}))
        if missing:
            notes.append(_not_checked("leakage", "key_column_missing", {**subject, "column": ",".join(missing)}, "The subject key column(s) %s are not in %s." % (", ".join(missing), table.name)))
            continue
        key_indexes = [table.columns.index(c) for c in keys]
        value_index, time_index = table.index(column), table.index(time_column)
        tokens = set(declared.get("tables", {}).get(table.name, {}).get("variables", {}).get(column, {}).get("missingTokens", ()))
        checked = leaked = no_cutoff = unparseable = 0
        leaked_rows, leaked_subjects = [], set()
        for row, number in zip(table.rows, table.numbers):
            if is_missing(row[value_index], tokens) or is_standard_missing(row[value_index]):
                continue
            key = tuple(row[i].strip() for i in key_indexes)
            limit = cutoffs.get(key)
            if limit is None:
                no_cutoff += 1
                continue
            measured = parse_time(row[time_index]) if row[time_index].strip() else None
            if measured is None:
                unparseable += 1
                continue
            checked += 1
            # A date with no time of day cannot be "after" a moment on the same day.
            after = measured[0].date() > limit[0].date() if not (measured[1] and limit[1]) else measured[0] > limit[0]
            if after:
                leaked += 1
                leaked_subjects.add(key)
                if len(leaked_rows) < SAMPLE_ROWS:
                    leaked_rows.append(number)
        if checked == 0:
            notes.append(_not_checked("leakage", "time_unparseable", subject, "No measurement of %s could be compared with a cutoff (%d without a cutoff, %d with an unreadable time)." % (column, no_cutoff, unparseable)))
            continue
        if leaked:
            found.append(finding("temporal_leakage", subject,
                                 "%d of %d measurements of %s (%d subject(s)) are after the cutoff %s; a model using %s would be reading the future." % (
                                     leaked, checked, column, len(leaked_subjects), cutoff["column"], column),
                                 count=leaked, rows=leaked_rows,
                                 detail={"rowsChecked": checked, "subjects": len(leaked_subjects), "cutoffColumn": cutoff["column"], "timeColumn": time_column,
                                         "rowsWithoutCutoff": no_cutoff, "rowsTimeUnreadable": unparseable}))
        else:
            clean.append(_clean("leakage", subject))


# --- transformations --------------------------------------------------------

def check_transformations(transformations: list, workspace: str, found: list, clean: list) -> None:
    for record in transformations:
        code = record.get("code")
        if not code:
            continue
        subject = {"step": record["name"]}
        try:
            current = hashlib.sha256(read_workspace_bytes(workspace, code["path"], 8 * 1024 * 1024)).hexdigest()
        except TableError:
            found.append(finding("transformation_code_missing", subject, "The code that made %s (v%d), %s, is not in the workspace." % (record["name"], record["version"], code["path"]),
                                 detail={"version": record["version"], "path": code["path"]}))
            continue
        if current != code["sha256"]:
            found.append(finding("transformation_code_changed", subject,
                                 "%s (v%d) was recorded from %s at %s; the file there now hashes %s. Re-record it to take the new version." % (
                                     record["name"], record["version"], code["path"], code["sha256"][:8], current[:8]),
                                 detail={"version": record["version"], "path": code["path"], "recorded": code["sha256"][:8], "current": current[:8]}))
        else:
            clean.append(_clean("transformations", subject))


# --- the whole check --------------------------------------------------------

def match_tables(tables: list, resolved: dict, explicit: dict):
    """{asset table name: Table} for the tables read, and the names that matched nothing.

    By the caller's explicit mapping, else by name, else by the one recorded table whose columns the
    file shares most (at least 60%), said as `matchedBy`.
    """
    matched, how, unmatched = {}, {}, []
    recorded = resolved["bindings"]
    for table in tables:
        target = explicit.get(table.name) or explicit.get(table.path)
        if target:
            matched[target] = table
            how[target] = "mapping"
        elif table.name in resolved["tables"] or table.name in recorded:
            matched[table.name] = table
            how[table.name] = "name"
        else:
            best, best_share = None, 0.0
            for name, binding in recorded.items():
                if name in matched or name in {t.name for t in tables}:
                    continue
                old = {c["name"] for c in binding.get("columns", [])}
                share = len(old & set(table.columns)) / len(old | set(table.columns)) if old | set(table.columns) else 0.0
                if share > best_share:
                    best, best_share = name, share
            if best and best_share >= 0.6:
                matched[best] = table
                how[best] = "columns:%.0f%%" % (best_share * 100)
            else:
                unmatched.append(table)
    return matched, how, unmatched


def run_checks(request: dict) -> dict:
    """One check of the files against the asset.

    `request`: {workspace, files: [{path, table?}], asset (the stored asset or None), complete? (the files are the
    whole delivery, so a recorded table with no file is `table_removed`), observationKeys?, subjectKeys?, joins?,
    leakage?, steps?}. Returns {findings, notChecked, clean, tables, denominators, joins, warnings}. A file
    nobody can read is reported as not checked, never raised: the rest of the files are still checked.
    """
    resolved = resolve(request.get("asset"), request)
    found, notes, clean, tables_out, warnings, profiles = [], [], [], [], [], []
    read, explicit = [], {}
    for entry in request["files"]:
        try:
            tables = read_tables(request["workspace"], entry["path"])
        except TableError as error:
            notes.append(_not_checked("drift", error.reason, {"table": PurePosixPath(entry["path"]).name}, str(error)))
            continue
        for table in tables:
            read.append(table)
            if entry.get("table") and len(tables) == 1:
                explicit[table.name] = entry["table"]
    matched, how, unmatched = match_tables(read, resolved, explicit)
    for table in unmatched:
        if request.get("asset"):
            found.append(finding("table_added", {"table": table.name}, "%s matches no table the dataset records (%d rows, %d columns)." % (table.name, len(table.rows), len(table.columns)),
                                 detail={"rows": len(table.rows), "columns": len(table.columns)}))
        matched.setdefault(table.name, table)
        how.setdefault(table.name, "new")
    # A table that is not among the files named is only "removed" when the caller says the files are the whole delivery.
    if request.get("asset") and request.get("complete"):
        for name in resolved["bindings"]:
            if name not in matched and name in resolved["tables"]:
                found.append(finding("table_removed", {"table": name}, "The recorded table %s has no file among those checked." % name))
    for name, table in matched.items():
        declared_table = resolved["tables"].get(name, {"subjectKey": None, "observationKey": None, "variables": {}, "identifierColumns": set()})
        tables_out.append({"table": name, "file": table.name, "path": table.path, "sha256": table.sha256, "bytes": table.bytes, "rows": len(table.rows),
                           "columns": len(table.columns), "matchedBy": how[name]})
        if request.get("asset"):
            check_drift(table, resolved["bindings"].get(name), declared_table, found, notes, clean)
        else:
            notes.append(_not_checked("drift", "no_asset", {"table": table.name}, "No meaning is recorded for this dataset yet, so there is nothing to compare %s with." % table.name))
        check_duplicates(table, declared_table, found, notes, clean, warnings)
    by_name = {table.name: table for table in read}
    by_name.update(matched)
    for join in resolved["joins"]:
        check_join(join, by_name, resolved, found, notes, clean, warnings, profiles)
    if request.get("leakage") is not None:
        check_leakage(request["leakage"], by_name, resolved, found, notes, clean)
    denominators = {}
    if request.get("steps"):
        steps = []
        for step in request["steps"]:
            step = dict(step)
            if step.get("rows") is None and step.get("path"):
                try:
                    opened = read_tables(request["workspace"], step["path"])
                    chosen = next((t for t in opened if t.name == step.get("table")), opened[0] if opened else None)
                except TableError as error:
                    notes.append(_not_checked("denominators", error.reason, {"step": str(step["label"])}, str(error)))
                    continue
                if chosen is not None:
                    by_name["step:" + str(step["label"])] = chosen
                    step["table"] = "step:" + str(step["label"])
            steps.append(step)
        denominators = check_denominators(steps, by_name, resolved, resolved["denominators"], found, notes, clean)
    check_transformations(resolved["transformations"], request["workspace"], found, clean)
    return {"findings": found, "notChecked": notes, "clean": clean, "tables": tables_out, "denominators": denominators, "joins": profiles, "warnings": warnings}


def file_sha256(workspace: str, relative: str, limit: int = 8 * 1024 * 1024) -> str:
    """The SHA-256 of one workspace file (a script, a dictionary): the identity a record names it by."""
    return hashlib.sha256(read_workspace_bytes(workspace, relative, limit)).hexdigest()
