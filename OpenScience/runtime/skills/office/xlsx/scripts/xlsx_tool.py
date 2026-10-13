#!/usr/bin/env python3
"""Look inside a workbook, and make its formulas show their values.

    xlsx_tool.py describe sales.xlsx
    xlsx_tool.py check budget.xlsx
    xlsx_tool.py recalc budget.xlsx --out budget-calculated.xlsx
    xlsx_tool.py csv sales.xlsx --sheet Data --range A1:E20

A workbook written by openpyxl or XlsxWriter carries formulas but no results: the numbers appear
only when an application recalculates. `recalc` has the LibreOffice Kit that ships with the
runtime refresh every formula and writes a new workbook whose cells hold their values (the
formulas stay formulas), and `check` reports formulas without a result and cells that hold an
error (#REF!, #DIV/0!, #NAME?, #VALUE!, #N/A). Run `recalc` last, then `check` the output.
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "shared"))
import office_kit  # noqa: E402

ERRORS = ("#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!")


def cell_text(value: object) -> str:
    return "" if value is None else str(value)


def scan(path: Path) -> dict:
    """Formulas, their cached results, and error cells, sheet by sheet."""
    formulas = openpyxl.load_workbook(path)
    values = openpyxl.load_workbook(path, data_only=True)
    report: dict = {"sheets": [], "formulas": 0, "without_result": [], "errors": []}
    for sheet in formulas.worksheets:
        result = values[sheet.title]
        count = 0
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    count += 1
                    cached = result[cell.coordinate].value
                    if cached is None:
                        report["without_result"].append(f"{sheet.title}!{cell.coordinate}")
                if isinstance(result[cell.coordinate].value, str) and result[cell.coordinate].value.strip() in ERRORS:
                    report["errors"].append({"cell": f"{sheet.title}!{cell.coordinate}", "value": result[cell.coordinate].value})
        report["formulas"] += count
        report["sheets"].append({"name": sheet.title, "dimensions": sheet.dimensions, "rows": sheet.max_row, "columns": sheet.max_column,
                                 "formulas": count, "merged": [str(r) for r in sheet.merged_cells.ranges][:10], "charts": len(sheet._charts), "images": len(sheet._images),
                                 "state": sheet.sheet_state})
    return report


def cmd_describe(args: argparse.Namespace) -> int:
    report = scan(args.file)
    book = openpyxl.load_workbook(args.file, data_only=True)
    for sheet in report["sheets"]:
        print(f"## {sheet['name']}  {sheet['dimensions']}  rows={sheet['rows']} columns={sheet['columns']} formulas={sheet['formulas']} charts={sheet['charts']} images={sheet['images']}")
        for row in book[sheet["name"]].iter_rows(min_row=1, max_row=args.rows, values_only=True):
            print(" | ".join(cell_text(v)[:24] for v in row[: args.columns]))
    print(f"formulas: {report['formulas']}, without a stored result: {len(report['without_result'])}, error cells: {len(report['errors'])}")
    return 0


def cmd_check(args: argparse.Namespace) -> int:
    report = scan(args.file)
    report["without_result"] = report["without_result"][:20]
    report["ok"] = not report["errors"] and not report["without_result"]
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["ok"] else 1


def cmd_recalc(args: argparse.Namespace) -> int:
    destination = args.out
    if destination.exists():
        print(json.dumps({"ok": False, "error": f"{destination} exists; choose a new file name"}))
        return 2
    try:
        engine = office_kit.recalculate(args.file, destination)
    except (office_kit.KitError, OSError) as error:
        print(json.dumps({"ok": False, "error": str(error)}, ensure_ascii=False))
        return 1
    report = scan(destination)
    report["ok"] = not report["errors"] and not report["without_result"]
    report["out"] = str(destination)
    report["missing_fonts"] = engine.get("missingFonts", [])
    report["without_result"] = report["without_result"][:20]
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["ok"] else 1


def cmd_csv(args: argparse.Namespace) -> int:
    book = openpyxl.load_workbook(args.file, data_only=True)
    sheet = book[args.sheet] if args.sheet else book.worksheets[0]
    rows = sheet[args.range] if args.range else sheet.iter_rows()
    writer = csv.writer(sys.stdout)
    for row in rows:
        writer.writerow([cell_text(c.value) for c in row])
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", required=True)
    describe = sub.add_parser("describe", help="sheets, sizes, first rows, formula and chart counts")
    describe.add_argument("file", type=Path)
    describe.add_argument("--rows", type=int, default=6)
    describe.add_argument("--columns", type=int, default=12)
    check = sub.add_parser("check", help="formulas without a result, cells in error")
    check.add_argument("file", type=Path)
    recalc = sub.add_parser("recalc", help="write a copy whose formulas carry their results")
    recalc.add_argument("file", type=Path)
    recalc.add_argument("--out", type=Path, required=True)
    csv_cmd = sub.add_parser("csv", help="print a sheet or a range as CSV (results, not formulas)")
    csv_cmd.add_argument("file", type=Path)
    csv_cmd.add_argument("--sheet")
    csv_cmd.add_argument("--range")
    args = parser.parse_args()
    return {"describe": cmd_describe, "check": cmd_check, "recalc": cmd_recalc, "csv": cmd_csv}[args.command](args)


if __name__ == "__main__":
    sys.exit(main())
