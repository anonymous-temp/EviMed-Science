#!/usr/bin/env python3
"""The Office skills' read / edit / create flows, run on real files with the pinned libraries.

    python3 test_office_flows.py            # skips what needs the LibreOffice Kit when it is not installed
    EVIMED_OFFICE_REQUIRE_KIT=1 python3 test_office_flows.py     # the image build: a missing kit is a failure

Each test is one thing a person asks an assistant to do to a document: change one clause and leave the rest
alone, add a column and a totals row with real formulas, split a slide and switch the deck to a dark palette
in the same file, pull a table out of a PDF. The assertions look at the files, not at the tools' words.
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))

import office_fixtures as fixtures  # noqa: E402
import office_kit  # noqa: E402

DOCX = ROOT / "docx/scripts/docx_tool.py"
XLSX = ROOT / "xlsx/scripts/xlsx_tool.py"
PPTX = ROOT / "pptx/scripts/pptx_tool.py"
PDF = ROOT / "pdf/scripts/pdf_tool.py"

try:
    office_kit.find_cli()
    KIT = shutil.which("node") is not None
except office_kit.KitError:
    KIT = False
if os.environ.get("EVIMED_OFFICE_REQUIRE_KIT") and not KIT:
    raise SystemExit("EVIMED_OFFICE_REQUIRE_KIT is set and the LibreOffice Kit CLI is not installed")
needs_kit = unittest.skipUnless(KIT, "the LibreOffice Kit CLI is not installed here")
try:
    office_kit.find_checker()
    CHECKER = True
except office_kit.KitError:
    CHECKER = False
needs_checker = unittest.skipUnless(CHECKER, "the kernel's Office checker is not installed here")


def run(script: Path, *arguments: object) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(script), *map(str, arguments)], capture_output=True, text=True, timeout=300)


def members(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


class Workdir(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory(prefix="evimed-office-")
        self.dir = Path(self._tmp.name)

    def tearDown(self) -> None:
        self._tmp.cleanup()


class DocxFlows(Workdir):
    def setUp(self) -> None:
        super().setUp()
        self.contract = fixtures.make_contract(self.dir / "contract.docx")

    def test_one_clause_changes_and_nothing_else_moves(self) -> None:
        from docx import Document

        edited = self.dir / "contract-v2.docx"
        done = run(DOCX, "replace", self.contract, "--find", "30 days", "--with", "45 days", "--paragraph", 18, "--out", edited)
        self.assertEqual(done.returncode, 0, done.stdout + done.stderr)
        before, after = members(self.contract), members(edited)
        self.assertEqual(sorted(before), sorted(after))
        changed = sorted(name for name in before if before[name] != after[name])
        self.assertEqual(changed, ["word/document.xml"], "only the body part may be rewritten; styles, header, footer and settings stay byte for byte")
        report = json.loads(run(DOCX, "diff", self.contract, edited).stdout)
        self.assertEqual(len(report["paragraph_changes"]), 1)
        self.assertEqual(report["paragraph_changes"][0]["before"], ["Payment is due within 30 days of the invoice date."])
        self.assertEqual(report["paragraph_changes"][0]["after"], ["Payment is due within 45 days of the invoice date."])
        # the second "30 days" (late fees) and the Chinese clause are untouched, and the bold run is still the only bold one
        document = Document(edited)
        texts = [p.text for p in document.paragraphs]
        self.assertIn("Late fees accrue after 30 days of delay at 1% per month.", texts)
        clause = next(p for p in document.paragraphs if p.text.startswith("Payment is due"))
        self.assertEqual([(r.text, bool(r.bold)) for r in clause.runs], [("Payment is due within ", False), ("45", True), (" days of the invoice date.", False)])
        self.assertEqual(document.sections[0].header.paragraphs[0].text, "Services Agreement - Confidential")
        self.assertEqual(len(Document(self.contract).paragraphs), len(document.paragraphs))

    def test_a_chinese_contract_changes_in_the_clause_named_and_not_in_its_twin(self) -> None:
        from docx import Document

        contract = fixtures.make_contract_zh(self.dir / "zh.docx")
        refused = run(DOCX, "replace", contract, "--find", "30", "--with", "45", "--out", self.dir / "never.docx")
        self.assertEqual(refused.returncode, 3, "the same number stands in the payment and the liability clause")
        edited = self.dir / "zh-45.docx"
        done = run(DOCX, "replace", contract, "--find", "30天", "--with", "45天", "--paragraph", 9, "--out", edited)
        self.assertEqual(done.returncode, 0, done.stdout)
        texts = [p.text for p in Document(edited).paragraphs]
        self.assertIn("\u7532\u65b9\u5e94\u81ea\u6536\u5230\u53d1\u7968\u4e4b\u65e5\u8d7745\u5929\u5185\u652f\u4ed8\u5168\u90e8\u6b3e\u9879\uff0c\u652f\u4ed8\u65b9\u5f0f\u4e3a\u94f6\u884c\u8f6c\u8d26\u3002", texts)
        self.assertIn("\u4e59\u65b9\u903e\u671f\u4ea4\u4ed8\u8d85\u8fc7 30 \u5929\u7684\uff0c\u7532\u65b9\u6709\u6743\u89e3\u9664\u5408\u540c\u5e76\u8981\u6c42\u8d54\u507f\u3002", texts)
        self.assertEqual(json.loads(run(DOCX, "diff", contract, edited).stdout)["parts_changed"], ["word/document.xml"])

    def test_an_ambiguous_change_is_refused_and_names_the_candidates(self) -> None:
        out = self.dir / "never.docx"
        done = run(DOCX, "replace", self.contract, "--find", "30 days", "--with", "45 days", "--out", out)
        self.assertEqual(done.returncode, 3)
        self.assertFalse(out.exists())
        self.assertEqual([m["paragraph"] for m in json.loads(done.stdout)["matches"]], [18, 19])

    def test_a_missing_text_is_reported_not_guessed(self) -> None:
        done = run(DOCX, "replace", self.contract, "--find", "sixty days", "--with", "45 days", "--out", self.dir / "x.docx")
        self.assertEqual(done.returncode, 2)

    def test_a_match_that_spans_runs_is_found_and_replaced(self) -> None:
        from docx import Document

        found = run(DOCX, "find", self.contract, "due within 30 days")
        self.assertEqual(found.returncode, 0, found.stdout)
        done = run(DOCX, "replace", self.contract, "--find", "due within 30 days of", "--with", "due within 45 days of", "--paragraph", 18)
        self.assertEqual(done.returncode, 0, done.stdout)
        self.assertTrue(any(p.text == "Payment is due within 45 days of the invoice date." for p in Document(self.contract).paragraphs), "no --out edits the file itself")

    def test_a_tracked_change_is_a_real_revision(self) -> None:
        from docx import Document

        edited = self.dir / "tracked.docx"
        done = run(DOCX, "replace", self.contract, "--find", "30 days", "--with", "45 days", "--paragraph", 18, "--track-changes", "--author", "Reviewer", "--out", edited)
        self.assertEqual(done.returncode, 0, done.stdout)
        body = members(edited)["word/document.xml"].decode("utf-8")
        self.assertRegex(body, r'<w:del [^>]*w:author="Reviewer"[^>]*><w:r><w:rPr><w:b/></w:rPr><w:delText[^>]*>30</w:delText></w:r></w:del>')
        self.assertRegex(body, r'<w:ins [^>]*w:author="Reviewer"[^>]*><w:r><w:rPr><w:b/></w:rPr><w:t[^>]*>45</w:t></w:r></w:ins>')
        Document(edited)  # still opens

    @needs_kit
    def test_the_edited_contract_still_lays_out_on_three_pages(self) -> None:
        edited = self.dir / "contract-v2.docx"
        self.assertEqual(run(DOCX, "replace", self.contract, "--find", "30 days", "--with", "45 days", "--paragraph", 18, "--out", edited).returncode, 0)
        for name in ("contract", "contract-v2"):
            office_kit.convert(self.dir / f"{name}.docx", self.dir / f"{name}.pdf")
        counts = [json.loads(run(PDF, "info", self.dir / f"{name}.pdf").stdout)["pages"] for name in ("contract", "contract-v2")]
        self.assertEqual(counts, [3, 3])


class KernelChecker(Workdir):
    @needs_checker
    def test_the_kernels_own_structural_check_reads_every_format(self) -> None:
        contract = fixtures.make_contract(self.dir / "contract.docx")
        report, status = office_kit.check(contract, contains=["Payment is due within 30 days"])
        self.assertEqual((status, report["verdict"], report["summary"]["paragraphs"]), (0, "pass", 25))
        self.assertEqual(office_kit.check(contract, contains=["a sentence the contract does not hold"])[1], 1)
        self.assertEqual(office_kit.check(fixtures.make_deck(self.dir / "deck.pptx"), count=10)[0]["summary"]["slides"], 10)
        sheets = office_kit.check(fixtures.make_sales(self.dir / "sales.xlsx", rows=20))[0]["summary"]["sheets"]
        self.assertEqual(sheets[0]["name"], "Data")


class XlsxFlows(Workdir):
    def setUp(self) -> None:
        super().setUp()
        self.sales = fixtures.make_sales(self.dir / "sales.xlsx")

    def add_column_and_totals(self) -> Path:
        import openpyxl

        book = openpyxl.load_workbook(self.sales)
        sheet = book["Data"]
        sheet["F1"] = "Revenue"
        last = sheet.max_row
        for row in range(2, last + 1):
            sheet[f"F{row}"] = f"=D{row}*E{row}"
        sheet[f"A{last + 1}"] = "Total"
        sheet[f"D{last + 1}"] = f"=SUM(D2:D{last})"
        sheet[f"F{last + 1}"] = f"=SUM(F2:F{last})"
        target = self.dir / "sales-with-formulas.xlsx"
        book.save(target)
        return target

    def test_describe_sees_all_rows(self) -> None:
        done = run(XLSX, "describe", self.sales)
        self.assertEqual(done.returncode, 0, done.stderr)
        self.assertIn("rows=5001", done.stdout)

    def test_a_workbook_with_formulas_but_no_results_is_reported(self) -> None:
        done = run(XLSX, "check", self.add_column_and_totals())
        self.assertEqual(done.returncode, 1)
        self.assertEqual(json.loads(done.stdout)["formulas"], 5002)
        self.assertTrue(json.loads(done.stdout)["without_result"])

    @needs_kit
    def test_formulas_are_recalculated_and_stay_formulas(self) -> None:
        import openpyxl

        source = self.add_column_and_totals()
        calculated = self.dir / "sales-calculated.xlsx"
        done = run(XLSX, "recalc", source, "--out", calculated)
        self.assertEqual(done.returncode, 0, done.stdout + done.stderr)
        report = json.loads(done.stdout)
        self.assertTrue(report["ok"])
        self.assertEqual(report["errors"], [])
        values = openpyxl.load_workbook(calculated, data_only=True)["Data"]
        self.assertAlmostEqual(values["F5002"].value, fixtures.sales_revenue(), places=2)
        self.assertEqual(openpyxl.load_workbook(calculated)["Data"]["F5002"].value, "=SUM(F2:F5001)")
        self.assertEqual(run(XLSX, "recalc", source, "--out", calculated).returncode, 2, "the kit never overwrites; neither does the tool")

    @needs_kit
    def test_an_error_cell_is_named(self) -> None:
        import openpyxl

        book = openpyxl.Workbook()
        book.active["A1"] = "=1/0"
        book.active["A2"] = "=SUM(B1:B3)"
        source = self.dir / "broken.xlsx"
        book.save(source)
        done = run(XLSX, "recalc", source, "--out", self.dir / "broken-calculated.xlsx")
        self.assertEqual(done.returncode, 1)
        self.assertEqual([e["cell"] for e in json.loads(done.stdout)["errors"]], ["Sheet!A1"])


class PptxFlows(Workdir):
    def setUp(self) -> None:
        super().setUp()
        self.deck = fixtures.make_deck(self.dir / "deck.pptx")

    def slide_texts(self, path: Path) -> list[list[str]]:
        from pptx import Presentation

        return [[s.text_frame.text for s in slide.shapes if s.has_text_frame] for slide in Presentation(path).slides]

    def test_a_crowded_slide_is_split_in_the_same_file(self) -> None:
        from pptx import Presentation

        before = self.slide_texts(self.deck)
        self.assertEqual(run(PPTX, "duplicate", self.deck, "--slide", 3).returncode, 0)
        self.assertEqual(run(PPTX, "delete-shapes", self.deck, "--slide", 3, "--ids", 4).returncode, 0)
        self.assertEqual(run(PPTX, "set-text", self.deck, "--slide", 3, "--shape", 3, "--text", "Point 3.1\\nPoint 3.2\\nPoint 3.3").returncode, 0)
        self.assertEqual(run(PPTX, "set-text", self.deck, "--slide", 4, "--shape", 2, "--text", "Section 3 (continued)").returncode, 0)
        deck = Presentation(self.deck)
        self.assertEqual(len(deck.slides), 11)
        after = self.slide_texts(self.deck)
        self.assertEqual(after[:2] + after[4:], before[:2] + before[3:], "every other slide keeps its content")
        self.assertEqual(after[2][1], "Point 3.1\nPoint 3.2\nPoint 3.3")
        self.assertEqual(after[3][0], "Section 3 (continued)")
        # the chart went with the second slide, and what the first slide no longer shows is not left behind in the file
        names = [str(shape.chart.part.partname) for slide in deck.slides for shape in slide.shapes if shape.has_chart]
        self.assertEqual(len(names), 1)
        self.assertEqual([n for n in members(self.deck) if n.startswith(("ppt/charts/chart", "ppt/embeddings/"))].__len__(), 2)

    def test_a_copied_slide_has_its_own_chart(self) -> None:
        from pptx import Presentation

        self.assertEqual(run(PPTX, "duplicate", self.deck, "--slide", 3).returncode, 0)
        deck = Presentation(self.deck)
        names = [str(shape.chart.part.partname) for slide in deck.slides for shape in slide.shapes if shape.has_chart]
        self.assertEqual(len(set(names)), 2, "editing one chart must not edit the other")
        self.assertEqual(len([n for n in members(self.deck) if n.startswith("ppt/embeddings/")]), 2)

    def test_slides_move_and_go(self) -> None:
        from pptx import Presentation

        self.assertEqual(run(PPTX, "move", self.deck, "--slide", 10, "--to", 2).returncode, 0)
        self.assertEqual(Presentation(self.deck).slides[1].shapes.title.text, "Section 10")
        self.assertEqual(run(PPTX, "delete", self.deck, "--slide", 2).returncode, 0)
        deck = Presentation(self.deck)
        self.assertEqual(len(deck.slides), 9)
        self.assertNotIn("Section 10", [s.shapes.title.text for s in deck.slides])

    def test_the_dark_palette_swaps_the_scheme_roles(self) -> None:
        out = self.dir / "dark.pptx"
        done = run(PPTX, "theme", self.deck, "--mode", "dark", "--out", out)
        self.assertEqual(done.returncode, 0, done.stdout)
        parts = members(out)
        master = next(v for k, v in parts.items() if k.startswith("ppt/slideMasters/slideMaster") and k.endswith(".xml")).decode()
        self.assertRegex(master, r'<p:clrMap [^>]*bg1="dk1"[^>]*tx1="lt1"')
        theme = next(v for k, v in parts.items() if k.startswith("ppt/theme/theme")).decode()
        self.assertRegex(theme, r'<a:dk1><a:srgbClr val="0F172A"/>')
        self.assertEqual(self.slide_texts(out), self.slide_texts(self.deck), "a palette change touches no text")

    @needs_kit
    def test_the_dark_deck_renders_dark_and_the_light_deck_renders_light(self) -> None:
        from PIL import Image

        out = self.dir / "dark.pptx"
        self.assertEqual(run(PPTX, "theme", self.deck, "--mode", "dark", "--out", out).returncode, 0)

        def brightness(deck: Path, name: str) -> float:
            manifest = office_kit.render(deck, self.dir / name, pages="3", dpi=40)
            pixels = list(Image.open(manifest["images"][0]["path"]).convert("L").getdata())
            return sum(pixels) / len(pixels) / 255

        self.assertLess(brightness(out, "dark-pages"), 0.3)
        self.assertGreater(brightness(self.deck, "light-pages"), 0.8)


class PdfFlows(Workdir):
    def setUp(self) -> None:
        super().setUp()
        self.report = fixtures.make_pdf(self.dir / "report.pdf")

    def test_info_and_find(self) -> None:
        info = json.loads(run(PDF, "info", self.report).stdout)
        self.assertEqual((info["pages"], info["pages_without_text_layer"]), (20, []))
        found = run(PDF, "find", self.report, "net revenue")
        self.assertIn("page 5:", found.stdout)
        self.assertTrue(found.stdout.strip().endswith("20 line(s)"))

    def test_a_table_leaves_the_pdf_cell_for_cell(self) -> None:
        import openpyxl

        out = self.dir / "table.xlsx"
        done = run(PDF, "tables", self.report, "--page", 5, "--xlsx", out)
        self.assertEqual(done.returncode, 0, done.stdout + done.stderr)
        rows = [[str(c) for c in row] for row in openpyxl.load_workbook(out).active.iter_rows(values_only=True)]
        self.assertEqual(rows, fixtures.PDF_TABLE)
        numbers = self.dir / "numbers.xlsx"
        self.assertEqual(run(PDF, "tables", self.report, "--page", 5, "--xlsx", numbers, "--numbers").returncode, 0)
        sheet = openpyxl.load_workbook(numbers).active
        self.assertEqual((sheet["B2"].value, sheet["B4"].value, sheet["A2"].value), (120, 45.5, "Alpha"))

    def test_a_page_without_a_table_says_so(self) -> None:
        self.assertEqual(run(PDF, "tables", self.report, "--page", 6).returncode, 2)

    @needs_kit
    def test_a_page_renders_to_an_image(self) -> None:
        done = run(PDF, "render", self.report, "--pages", 5, "--out-dir", self.dir / "pages", "--dpi", 50)
        self.assertEqual(done.returncode, 0, done.stdout + done.stderr)
        self.assertTrue(Path(json.loads(done.stdout)["pages"][0]).stat().st_size > 1000)


if __name__ == "__main__":
    unittest.main(verbosity=2)
