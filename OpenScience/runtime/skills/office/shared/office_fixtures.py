"""Small, deterministic fixtures for the Office flows: a 3-page contract, a 5,000-row sales workbook,
a 20-page PDF with a ruled table on page 5, and a 10-slide deck with a chart on slide 3.

They are built with the same libraries the skills use (python-docx, XlsxWriter, python-pptx) and a few lines of raw PDF,
so a test needs neither network nor a committed binary.
"""
from __future__ import annotations

from pathlib import Path

PDF_TABLE = [["Product", "Q1", "Q2", "Q3"], ["Alpha", "120", "135", "150"], ["Beta", "80", "95", "70"],
             ["Gamma", "45.5", "50.25", "61"], ["Delta", "300", "280", "310"], ["Total", "545.5", "560.25", "591"]]


def make_contract(path: Path) -> Path:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.core_properties.title = "Services Agreement"
    document.core_properties.author = "Fixture"
    section = document.sections[0]
    section.header.paragraphs[0].text = "Services Agreement - Confidential"
    section.footer.paragraphs[0].text = "Page footer"
    document.add_heading("Services Agreement", 0)
    document.add_paragraph("This agreement is made between Northwind Ltd (the Client) and Contoso GmbH (the Supplier).")
    document.add_heading("1. Scope", 1)
    document.add_paragraph("The Supplier delivers the monthly reporting service described in Annex A.")
    document.add_page_break()
    document.add_heading("2. Fees", 1)
    paragraph = document.add_paragraph("Fees are set out in the table below and exclude VAT.")
    paragraph.runs[0].font.size = Pt(11)
    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for r, row in enumerate([["Item", "Unit", "Fee"], ["Reporting", "month", "1200"], ["Support", "hour", "95"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.add_heading("3. Payment", 1)
    clause = document.add_paragraph("Payment is due within ")
    clause.add_run("30").bold = True
    clause.add_run(" days of the invoice date.")
    document.add_paragraph("Late fees accrue after 30 days of delay at 1% per month.")
    document.add_paragraph("付款期限为收到发票后30天。")
    document.add_page_break()
    document.add_heading("4. Term", 1)
    document.add_paragraph("The agreement runs for twelve months and renews once for twelve months unless terminated.")
    document.add_heading("5. Signatures", 1)
    document.add_paragraph("Signed for the Client: ____________    Signed for the Supplier: ____________")
    document.save(path)
    return path


def make_sales(path: Path, rows: int = 5000) -> Path:
    import xlsxwriter

    book = xlsxwriter.Workbook(str(path))
    sheet = book.add_worksheet("Data")
    sheet.write_row(0, 0, ["Date", "Region", "Product", "Units", "Unit price"])
    regions = ["North", "South", "East", "West"]
    for i in range(rows):
        sheet.write_row(i + 1, 0, [f"2026-{(i % 12) + 1:02d}-{(i % 27) + 1:02d}", regions[i % 4], f"P{i % 17}", (i * 7) % 50 + 1, 10 + (i % 13) * 2.5])
    book.close()
    return path


def sales_revenue(rows: int = 5000) -> float:
    return sum(((i * 7) % 50 + 1) * (10 + (i % 13) * 2.5) for i in range(rows))


def make_pdf(path: Path, pages: int = 20, table_page: int = 5) -> Path:
    """A valid PDF written by hand: Helvetica text on every page, a ruled table on `table_page`."""
    objects: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for number in range(1, pages + 1):
        lines = [f"BT /F1 16 Tf 72 740 Td (Annual report - section {number}) Tj ET",
                 f"BT /F1 11 Tf 72 700 Td (Page {number} of {pages}. Net revenue grew in every region.) Tj ET"]
        if number == table_page:
            left, top, width, height = 72, 650, 90, 24
            for r in range(len(PDF_TABLE) + 1):
                y = top - r * height
                lines.append(f"{left} {y} m {left + width * 4} {y} l S")
            for c in range(5):
                x = left + c * width
                lines.append(f"{x} {top} m {x} {top - height * len(PDF_TABLE)} l S")
            for r, row in enumerate(PDF_TABLE):
                for c, value in enumerate(row):
                    lines.append(f"BT /F1 11 Tf {left + c * width + 6} {top - r * height - 16} Td ({value}) Tj ET")
        stream = "\n".join(lines).encode("latin-1")
        content = len(objects) + 2
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents {content} 0 R >>".encode())
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
        kids.append(f"{len(objects) - 1} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {pages} >>".encode()
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for index, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return path


def make_deck(path: Path, slides: int = 10) -> Path:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for number in range(1, slides + 1):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = f"Section {number}"
        box = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(5.6), Inches(4.5))
        frame = box.text_frame
        frame.word_wrap = True
        for line in range(1, 7 if number == 3 else 3):
            paragraph = frame.paragraphs[0] if line == 1 else frame.add_paragraph()
            paragraph.text = f"Point {number}.{line}"
            paragraph.runs[0].font.size = Pt(22)
        if number == 3:
            data = CategoryChartData()
            data.categories = ["Q1", "Q2", "Q3"]
            data.add_series("Revenue", [12, 18, 25])
            slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.8), Inches(1.6), Inches(5.8), Inches(4.5), data)
        slide.notes_slide.notes_text_frame.text = f"Notes for section {number}"
    deck.save(path)
    return path
