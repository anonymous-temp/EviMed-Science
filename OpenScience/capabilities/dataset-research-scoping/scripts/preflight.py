#!/usr/bin/env python3
"""Check scoped privacy and recomputable profiling; preserve partial research.

Only source identifier leakage is an artifact repair. Missing profiling inputs
or a failed recomputation are notices. Scientific feasibility and evidence
sufficiency belong to the analyst, not prose patterns or corpus quotas.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import subprocess
import sys
from datetime import date
import tempfile
from pathlib import Path

def read_text(root: Path, name: str) -> str:
    path = root / name
    return path.read_text(encoding="utf-8") if path.is_file() else ""


# Columns whose values identify a person or an episode of care. Their values may
# never leave the data; a run refers to subjects by pseudonyms it assigns.
# Kept identical to profile_dataset.py's rule. Matching the subject word alone
# swept in RECORD_DATE and RECORD_CONTENT, whose values are timestamps and
# clinical text: every date in the report would then have read as a leaked
# identifier and blocked a sound package.
# Matching the subject word as a substring was too broad (RECORD_DATE), and
# matching it as a whole word was too narrow: a real front-page column is
# MED_REC_NO — the medical record number, abbreviated — which neither rule
# caught, so its full vocabulary of hospital numbers was printed into the
# profile and handed over as a clean deliverable. Names are split into tokens
# and a column identifies when a subject token meets an id-shaped last token.
IDENTIFIER_EXPLICIT = re.compile(r"^(?:id|mrn|姓名|患者姓名|身份证号?|病案号|住院号|门诊号|就诊号)$", re.I)
IDENTIFIER_SUBJECT_TOKENS = frozenset({
    "patient", "subject", "person", "case", "record", "rec", "admission",
    "visit", "encounter", "inpatient", "outpatient", "mrn",
})
# "adm" and "reg" were in this set and masked ADM_DEPT_CODE, the admitting
# department — a covariate, not a person. A subject token must be a word that
# names the subject, not any abbreviation that begins one.
IDENTIFIER_SUFFIX_TOKENS = frozenset({"id", "no", "num", "number", "code", "sn"})
IDENTIFIER_SUBJECT_CJK = re.compile(r"(?:病案|住院|门诊|患者|病人|就诊|身份证|姓名)")


def is_identifying(name: str) -> bool:
    """True when a column's name says its values identify a person or an episode.

    Name only. `identifying_reason` below is what callers should ask: a column
    identifies when its name says so **or** when its values are shaped like an
    identity number, a phone number, an e-mail address or a date of birth.
    """
    cleaned = name.strip()
    if IDENTIFIER_EXPLICIT.match(cleaned):
        return True
    if IDENTIFIER_PERSONAL_NAMES.match(cleaned):
        return True
    tokens = [t for t in re.split(r"[^0-9A-Za-z]+", cleaned) if t]
    if tokens and tokens[-1].lower() in IDENTIFIER_SUFFIX_TOKENS:
        if any(t.lower() in IDENTIFIER_SUBJECT_TOKENS for t in tokens):
            return True
        if IDENTIFIER_SUBJECT_CJK.search(cleaned):
            return True
    return False


# A column name is not the only thing that identifies a person, and on real
# extracts it was not even the usual one. PATIENT_NAME, NAME, PHONE, TEL,
# ADDRESS, DOB, BIRTH_DATE, ID_CARD, 出生日期, 电话 and 住址 all read as ordinary
# columns to the rule above — none carries an id-shaped suffix token — so a
# twenty-row cohort had every personal name, every eighteen-digit identity
# number, every mobile number, every date of birth, every e-mail address and
# every home address printed into `data-profile.json` and `data-profile.md` as
# "vocabulary", and the preflight's leakage scan, which asks the same function,
# let all of it through. Reproduced on exactly that file: eight columns, six of
# them emitted raw.
#
# Two answers, and the second is the one that keeps working when a column is
# called `col_7`. First the names above, as names. Then the values themselves:
# an identity number that verifies its own check digit, a mainland mobile
# number, an e-mail address, a column of dates that are lifespans rather than
# events. These are closed format checks over values — arithmetic and fixed
# shapes — not judgements about language, which is the line development
# principle 1 draws and principle 5 forbids crossing.
IDENTIFIER_PERSONAL_NAMES = re.compile(
    r"^(?:"
    r"name|full[_\s-]?name|first[_\s-]?name|last[_\s-]?name|surname|given[_\s-]?name|patient[_\s-]?name"
    r"|phone|phone[_\s-]?no|tel|telephone|mobile|cell|contact[_\s-]?no"
    r"|address|addr|home[_\s-]?address|postcode|zip|zipcode"
    r"|dob|birth|birthday|birth[_\s-]?date|date[_\s-]?of[_\s-]?birth"
    r"|email|e[_\s-]?mail|mail"
    r"|id[_\s-]?card|idcard|id[_\s-]?no|idno|ssn|passport|passport[_\s-]?no"
    r"|姓名|名字|电话|手机|手机号码?|联系电话|联系方式|住址|地址|家庭住址|通讯地址|邮编"
    r"|出生日期|出生年月|生日|邮箱|电子邮箱|电子邮件|身份证|证件号码?|护照号码?"
    r")$",
    re.I,
)

MAINLAND_ID_18 = re.compile(r"^\d{17}[\dXx]$")
MAINLAND_ID_15 = re.compile(r"^\d{15}$")
# +86 or a bare 11-digit number beginning 13-19. Separators are stripped first.
MAINLAND_MOBILE = re.compile(r"^(?:\+?86)?1[3-9]\d{9}$")
EMAIL_ADDRESS = re.compile(r"^[^@\s]+@[^@\s.]+(?:\.[^@\s.]+)+$")
# ISO 7064:1983 MOD 11-2, the check digit every mainland identity number carries.
ID_CHECK_WEIGHTS = (7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2)
ID_CHECK_CODES = "10X98765432"
# A shape has to hold for most of the column, not for one cell: one stray
# "13800138000" in a free-text note must not mask a clinical column, and a
# column of identity numbers with two typos must still be masked. And a shape
# read off two values is a coincidence, so a short column is judged by name only.
VALUE_SHAPE_MIN_VALUES = 5
VALUE_SHAPE_MIN_SHARE = 0.8
# Nobody alive was born more than this long ago, and a birth is never in the
# future — the two bounds that make a lifespan distinguishable from an event.
BIRTH_YEAR_REACH = 130
# Three properties together separate a column of births from a column of events,
# and all three are needed. It spreads across at least this many years — a
# cohort contains people of different ages. Nothing in it is from the last year.
# And it reaches back at least a generation: every cohort has someone over
# thirty in it, while an event column, however long the follow-up, starts when
# the study did. A 2010-2016 admission window is narrow, a 2015-2026 one is
# current, and a twenty-year 2005-2025 follow-up — which passes both of the
# first two — is caught by the third.
BIRTH_YEAR_SPREAD = 15
BIRTH_YEAR_OLDEST_AT_LEAST = 30


def is_mainland_id_number(value: str) -> bool:
    """An 18-digit identity number that verifies its own check digit.

    The checksum is what makes this safe to run over every column in a hospital
    extract: an arbitrary 18-character code passes it about one time in eleven,
    a real identity number always. A 15-digit legacy number carries no check
    digit, so it is recognised only by its length and its date part.
    """
    cleaned = value.strip().replace(" ", "")
    if MAINLAND_ID_15.match(cleaned):
        return _plausible_id_birth(cleaned[6:12], century="19")
    if not MAINLAND_ID_18.match(cleaned):
        return False
    if not _plausible_id_birth(cleaned[6:14]):
        return False
    total = sum(int(cleaned[index]) * ID_CHECK_WEIGHTS[index] for index in range(17))
    return cleaned[17].upper() == ID_CHECK_CODES[total % 11]


def _plausible_id_birth(digits: str, century: str = "") -> bool:
    text = century + digits
    if len(text) != 8:
        return False
    year, month, day = int(text[:4]), int(text[4:6]), int(text[6:8])
    return 1 <= month <= 12 and 1 <= day <= 31 and 1900 <= year <= date.today().year


def is_mobile_number(value: str) -> bool:
    return bool(MAINLAND_MOBILE.match(re.sub(r"[\s()-]", "", value.strip())))


def is_email_address(value: str) -> bool:
    return bool(EMAIL_ADDRESS.match(value.strip()))


# A date, in either component order, with an optional time. Declared inside
# this block rather than borrowed, so the block is byte-identical in
# profile_dataset.py and preflight.py — the four-copy rule applies to both.
BIRTH_DATE_SHAPE = re.compile(
    r"^(?:\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}[-/]\d{1,2}[-/]\d{4})"
    r"(?:[ tT]\d{1,2}:\d{2}(?::\d{2})?)?\s*$",
)


def _parsed_year(value: str):
    if not BIRTH_DATE_SHAPE.match(value):
        return None
    head = re.split(r"[ tT]", value.strip())[0]
    parts = re.split(r"[-/]", head)
    for part in parts:
        if len(part) == 4 and part.isdigit():
            return int(part)
    return None


def looks_like_birth_dates(values: list) -> bool:
    """Dates that describe lifespans rather than events.

    A date column is not identifying because it is a date — admission dates,
    sample dates and follow-up dates are covariates and have to stay readable.
    What a column of birth dates has and an event column does not is reach: its
    values span decades, none of them is recent, and all of them are inside a
    human lifetime. Judged on the values, so a column called `col_7` is caught
    and `ADMISSION_DATE` is not.
    """
    years = [year for year in (_parsed_year(v) for v in values) if year is not None]
    if len(years) < VALUE_SHAPE_MIN_VALUES or len(years) < VALUE_SHAPE_MIN_SHARE * len(values):
        return False
    today = date.today()
    if min(years) < today.year - BIRTH_YEAR_REACH or max(years) > today.year:
        return False
    if max(years) > today.year - 1:
        return False
    if min(years) > today.year - BIRTH_YEAR_OLDEST_AT_LEAST:
        return False
    return max(years) - min(years) >= BIRTH_YEAR_SPREAD


def identifying_value_shape(values: list):
    """The shape that makes a column's values identify a person, or None.

    Returns the name of the shape so the profile can say why a column was
    masked. A masking nobody can see the reason for is one a researcher works
    around by exporting the column again under another name.

    Four shapes, and the list stops where format stops. A personal name and a
    street address have no closed format — recognising them means a pattern over
    language, which is the thing development principle 5 forbids and the thing
    that never converges. They are caught by column name, which is how they
    arrive in every real extract seen so far; an anonymised column of names is a
    known gap, recorded here rather than papered over with a surname list that
    would be wrong for most of the world.
    """
    candidates = [str(v).strip() for v in values if str(v).strip()]
    if len(candidates) < VALUE_SHAPE_MIN_VALUES:
        return None
    floor = VALUE_SHAPE_MIN_SHARE * len(candidates)
    for shape, matches in (
        ("id-number", is_mainland_id_number),
        ("mobile-number", is_mobile_number),
        ("email-address", is_email_address),
    ):
        if sum(1 for value in candidates if matches(value)) >= floor:
            return shape
    if looks_like_birth_dates(candidates):
        return "birth-date"
    return None


def identifying_reason(name: str, values: list):
    """Why a column is masked — its name, or the shape of what it holds."""
    if is_identifying(name):
        return "column-name"
    return identifying_value_shape(values)


# An identifier short enough to collide with an ordinary number in prose is not
# evidence of leakage; below this length a match is not reported.
IDENTIFIER_MIN_LENGTH = 5
def identifier_values(dataset_paths: list[Path]) -> set[str]:
    """Every value held by an identifier-shaped column of the source data."""
    values: set[str] = set()
    for path in dataset_paths:
        suffix = path.suffix.lower()
        tables: list[tuple[list[str], list[list[str]]]] = []
        if suffix in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook
            except ImportError:  # pragma: no cover - environment dependent
                continue
            workbook = load_workbook(path, read_only=True, data_only=True)
            for sheet in workbook.worksheets:
                rows = [["" if c is None else str(c) for c in row] for row in sheet.iter_rows(values_only=True)]
                if rows:
                    tables.append((rows[0], rows[1:]))
            workbook.close()
        else:
            delimiter = "\t" if suffix in {".tsv", ".tab"} else ","
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle, delimiter=delimiter))
            if rows:
                tables.append(([str(h) for h in rows[0]], [[str(c) for c in r] for r in rows[1:]]))

        for header, rows in tables:
            for index, column_name in enumerate(header):
                cells = [str(row[index]).strip() for row in rows if index < len(row)]
                # Name **or** value shape, the same rule the profiler applies.
                # Asking only the name meant a column of identity numbers called
                # ID_CARD contributed nothing to this set, so the deliverable
                # could carry every one of them and the leakage scan reported
                # clean — the profile and the preflight agreeing on the wrong
                # answer, which is worse than disagreeing.
                if identifying_reason(str(column_name), cells) is None:
                    continue
                for value in cells:
                    if len(value) >= IDENTIFIER_MIN_LENGTH:
                        values.add(value)
    return values


def scannable_files(root: Path, dataset_paths: list[Path]) -> list[Path]:
    """Every text file the run leaves behind, except the source data itself.

    Scanning only the declared deliverables was not enough. A run wrote its own
    working file holding {"pseudonyms": {"900004": "P1", ...}} — the mapping back
    to real hospital numbers, which defeats the pseudonyms entirely — and because
    that file was not on the declared list it was never looked at. Anything left
    in the workspace can be read by whoever receives it.
    """
    sources = {path.resolve() for path in dataset_paths}
    files = []
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.is_symlink():
            continue
        if path.resolve() in sources:
            continue
        # Retrieved source documents are quoted evidence, not run output.
        if ".evimed-sources" in path.parts:
            continue
        if path.suffix.lower() in {".xlsx", ".xlsm", ".xls", ".png", ".jpg", ".pdf", ".zip", ".gz"}:
            continue
        files.append(path)
    return files


def check_identifier_leakage(root: Path, dataset_paths: list[Path], issues: list[str]) -> int:
    values = identifier_values(dataset_paths)
    if not values:
        return 0
    leaked = 0
    for path in scannable_files(root, dataset_paths):
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
        except OSError:
            continue
        for value in sorted(values):
            if value in text:
                leaked += 1
                name = path.relative_to(root).as_posix()
                issues.append(
                    f"{name}: carries the source identifier from an identifying column. Redact this artifact; refer to subjects by a pseudonym you "
                    "assign, and never write the mapping back to the source value into the workspace."
                )
    return leaked


def check_profile_recomputable(root: Path, dataset_paths: list[Path], issues: list[str]) -> bool:
    script = root / "data-profile.py"
    recorded = root / "data-profile.json"
    if not script.is_file():
        issues.append("data-profile.py is missing: the profile numbers must come from a script that is kept.")
        return False
    if not recorded.is_file():
        issues.append("data-profile.json is missing: data-profile.py must write the machine-readable profile it renders.")
        return False
    if not dataset_paths:
        issues.append(
            "scoping-run.json: priorDataContact.filesReceived names no readable dataset file, "
            "so the profile cannot be recomputed."
        )
        return False
    with tempfile.TemporaryDirectory() as work:
        target = Path(work) / "recomputed.json"
        try:
            completed = subprocess.run(
                [sys.executable, str(script), *[str(p) for p in dataset_paths],
                 "--json", str(target), "--markdown", str(Path(work) / "recomputed.md")],
                cwd=root, capture_output=True, text=True, timeout=600,
            )
        except (OSError, subprocess.SubprocessError) as error:
            issues.append(f"data-profile.py could not be re-run: {error}")
            return False
        if completed.returncode != 0:
            issues.append(f"data-profile.py exited {completed.returncode} when re-run: {completed.stderr.strip()[:400]}")
            return False
        try:
            fresh = json.loads(target.read_text(encoding="utf-8"))
            stored = json.loads(recorded.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as error:
            issues.append(f"the recomputed profile could not be compared: {error}")
            return False
    if fresh != stored:
        issues.append(
            "data-profile.json does not match what data-profile.py produces from the same inputs; "
            "the reported numbers are not the script's."
        )
        return False
    return True


def dataset_inputs(root: Path, supplied: list[str], warnings: list[str]) -> list[Path]:
    entries = supplied
    if not entries:
        try:
            receipt = json.loads(read_text(root, "scoping-run.json") or "{}")
            contact = receipt.get("priorDataContact", {}) if isinstance(receipt, dict) else {}
            entries = contact.get("filesReceived", []) if isinstance(contact, dict) else []
        except (ValueError, OSError):
            warnings.append("scoping-run.json could not be read; pass --input for the source files.")
            entries = []
    if not isinstance(entries, list):
        entries = []
    paths = []
    for entry in entries:
        candidate = (root / str(entry)).resolve()
        if not candidate.is_relative_to(root) or not candidate.is_file():
            warnings.append("A dataset path is unavailable or outside the workspace; that file was not inspected.")
        else:
            paths.append(candidate)
    if not paths:
        warnings.append("No readable source files were supplied; privacy and recomputation could not be verified. Pass --input without creating a prior-contact receipt.")
    return sorted(set(paths))


def main() -> int:
    parser = argparse.ArgumentParser(description="Check dataset research scoping artifacts.")
    parser.add_argument("--workspace", default=".")
    parser.add_argument("--input", action="append", default=[])
    args = parser.parse_args()
    root = Path(args.workspace).resolve()
    issues: list[str] = []
    warnings: list[str] = []
    dataset_paths = dataset_inputs(root, args.input, warnings)
    leaked = check_identifier_leakage(root, dataset_paths, issues)
    recomputable = check_profile_recomputable(root, dataset_paths, warnings)
    payload = {
        "ok": not issues,
        "workspace": str(root),
        "metrics": {"datasetFiles": len(dataset_paths), "identifierLeaks": leaked,
                    "profileRecomputable": recomputable},
        "issues": issues,
        "warnings": warnings,
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0 if not issues else 1


if __name__ == "__main__":
    sys.exit(main())
